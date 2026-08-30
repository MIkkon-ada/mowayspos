"""Standalone, opt-in DeepSeek probes for spreadsheet understanding."""

from __future__ import annotations

import json
import time
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Callable, Literal

from openpyxl import load_workbook
from openpyxl.utils import quote_sheetname
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from .project_init_file_parser import parse_project_init_file


@dataclass(frozen=True)
class WorkbookEvidence:
    """Bounded workbook evidence that keeps exact worksheet locations."""

    text: str
    locations: set[str]
    merged_ranges: dict[str, str]
    input_sha256: str


class ProbeKeyTask(BaseModel):
    """One model-extracted key task accepted by the probe."""

    model_config = ConfigDict(extra="forbid")

    title: str
    owner: str = ""
    collaborators: list[str] = Field(default_factory=list)
    plan_start: str = ""
    plan_end: str = ""
    completion_standard: str = ""
    note: str = ""
    evidence: list[str] = Field(min_length=1)


class ProbeWorkstream(BaseModel):
    """One model-extracted workstream accepted by the probe."""

    model_config = ConfigDict(extra="forbid")

    title: str
    key_tasks: list[ProbeKeyTask] = Field(min_length=1)


class ProbeOutput(BaseModel):
    """Strict JSON envelope required from every provider route."""

    model_config = ConfigDict(extra="forbid")

    workstreams: list[ProbeWorkstream] = Field(min_length=1)


TEXT_MODELS = ("deepseek-v4-flash", "deepseek-v4-pro")
ProbeStatus = Literal[
    "succeeded",
    "invalid_json",
    "invalid_schema",
    "uncited_output",
    "upstream_error",
    "skipped",
]


@dataclass(frozen=True)
class ProbeResult:
    """A sanitized outcome for one probe model invocation."""

    mode: Literal["a", "b"]
    model_name: str
    status: ProbeStatus
    duration_ms: int
    output: ProbeOutput | None = None
    reason: str = ""


class ProbeRunner:
    """Compare text models against the same workbook evidence."""

    def __init__(self, *, complete_text: Callable[[str, str], str]) -> None:
        self._complete_text = complete_text

    def run_text(self, evidence: WorkbookEvidence) -> list[ProbeResult]:
        return [self._run_text_model(model_name, evidence) for model_name in TEXT_MODELS]

    def _run_text_model(self, model_name: str, evidence: WorkbookEvidence) -> ProbeResult:
        started = time.monotonic()
        try:
            response = self._complete_text(model_name, _text_prompt(evidence))
        except Exception:
            return ProbeResult(
                mode="a",
                model_name=model_name,
                status="upstream_error",
                duration_ms=_elapsed_ms(started),
            )
        try:
            payload = json.loads(response)
        except (TypeError, json.JSONDecodeError):
            return ProbeResult(
                mode="a",
                model_name=model_name,
                status="invalid_json",
                duration_ms=_elapsed_ms(started),
            )
        try:
            output = ProbeOutput.model_validate(payload)
        except ValidationError:
            return ProbeResult(
                mode="a",
                model_name=model_name,
                status="invalid_schema",
                duration_ms=_elapsed_ms(started),
            )
        if not _has_only_known_evidence(output, evidence.locations):
            return ProbeResult(
                mode="a",
                model_name=model_name,
                status="uncited_output",
                duration_ms=_elapsed_ms(started),
            )
        return ProbeResult(
            mode="a",
            model_name=model_name,
            status="succeeded",
            duration_ms=_elapsed_ms(started),
            output=output,
        )


def _elapsed_ms(started: float) -> int:
    return max(0, int((time.monotonic() - started) * 1000))


def _has_only_known_evidence(output: ProbeOutput, locations: set[str]) -> bool:
    return all(
        location in locations
        for workstream in output.workstreams
        for key_task in workstream.key_tasks
        for location in key_task.evidence
    )


def _text_prompt(evidence: WorkbookEvidence) -> str:
    return """你正在测试对 Excel 工作推进表的结构理解。只返回一个 JSON 对象，不能使用 Markdown。

请提取：专项/重点工作 -> workstreams；关键任务 -> key_tasks。每个关键任务必须包含
title、owner、collaborators、plan_start、plan_end、completion_standard、note、evidence。
“协同成员”“协助人”“参与人”以及类似“协助人：张三、李四”的备注都属于 collaborators；
不要在 note 中重复协助人名单。信息不确定时留空或在 note 说明，不能编造。
每个 key_task 必须用 evidence 引用下面证据中完全一致的单元格位置。

JSON 结构：
{"workstreams":[{"title":"","key_tasks":[{"title":"","owner":"","collaborators":[],"plan_start":"","plan_end":"","completion_standard":"","note":"","evidence":["'工作表'!A1"]}]}]}

证据：
""" + evidence.text


def build_workbook_evidence(
    path: Path,
    original_name: str,
    *,
    max_sheets: int,
    max_cells: int,
) -> WorkbookEvidence:
    """Return visible, non-empty cells and merge ranges after parser validation."""
    if max_sheets < 1 or max_cells < 1:
        raise ValueError("max_sheets and max_cells must be positive")

    input_path = Path(path)
    parsed_chunks = parse_project_init_file(input_path, original_name)
    workbook = load_workbook(input_path, read_only=False, data_only=True)
    try:
        lines = [
            f"[解析文本] {chunk.location}\n{chunk.text}"
            for chunk in parsed_chunks
            if chunk.text.strip()
        ]
        locations: set[str] = set()
        merged_ranges: dict[str, str] = {}
        emitted_cells = 0
        inspected_sheets = 0

        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            inspected_sheets += 1
            if inspected_sheets > max_sheets:
                break
            merged_anchors = {
                (merged.min_row, merged.min_col): str(merged)
                for merged in sheet.merged_cells.ranges
            }
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None or str(cell.value).strip() == "":
                        continue
                    emitted_cells += 1
                    if emitted_cells > max_cells:
                        raise ValueError("workbook evidence exceeds max_cells")
                    cell_range = merged_anchors.get((cell.row, cell.column), cell.coordinate)
                    location = f"{quote_sheetname(sheet.title)}!{cell_range}"
                    rendered_value = str(cell.value).strip()
                    locations.add(location)
                    lines.append(f"{location}={rendered_value}")
                    if cell_range != cell.coordinate:
                        merged_ranges[location] = rendered_value

        return WorkbookEvidence(
            text="\n".join(lines),
            locations=locations,
            merged_ranges=merged_ranges,
            input_sha256=sha256(input_path.read_bytes()).hexdigest(),
        )
    finally:
        workbook.close()
