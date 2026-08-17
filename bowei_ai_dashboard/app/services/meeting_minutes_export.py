"""Generate a downloadable DOCX meeting-minutes record from a saved draft."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from docx import Document

from .. import models


def _json_list(value: str | None) -> list[Any]:
    try:
        parsed = json.loads(value or "[]")
    except (TypeError, ValueError):
        return []
    return parsed if isinstance(parsed, list) else []


def _add_list(document: Document, title: str, values: list[Any]) -> None:
    document.add_heading(title, level=2)
    if not values:
        document.add_paragraph("（无）")
        return
    for value in values:
        if isinstance(value, dict):
            text = "；".join(f"{key}：{item}" for key, item in value.items() if item not in (None, ""))
        else:
            text = str(value)
        if text.strip():
            document.add_paragraph(text, style="List Bullet")


def build_meeting_minutes_docx(meeting: models.Meeting, *, source_name: str = "") -> bytes:
    document = Document()
    document.add_heading(meeting.title or "项目会议纪要", level=1)
    table = document.add_table(rows=0, cols=2)
    for label, value in (
        ("项目", meeting.related_special_project or ""),
        ("会议日期", meeting.meeting_date or ""),
        ("会议类型", meeting.meeting_type or ""),
        ("主持人", meeting.host or ""),
        ("参会人员", meeting.participants or ""),
        ("整理人", meeting.organizer or ""),
        ("原始文档", source_name),
        ("审核状态", meeting.review_status or meeting.publish_status or ""),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("会议小结", level=2)
    document.add_paragraph(meeting.summary or "（无）")
    _add_list(document, "会议行动项", _json_list(meeting.task_list_json))
    _add_list(document, "会议决议", _json_list(meeting.decision_items_json))
    _add_list(document, "风险与待确认事项", _json_list(meeting.risk_items_json))
    document.add_paragraph("本文件由项目会议纪要流程生成，执行安排变更须以项目负责人审核结果为准。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
