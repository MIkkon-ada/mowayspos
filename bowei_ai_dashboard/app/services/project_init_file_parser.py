"""Parse project-init source files into text chunks with source locations."""

from __future__ import annotations

import io
import subprocess
import threading
import zipfile
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import xlrd
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, quote_sheetname
from pypdf import PdfReader


MAX_INPUT_BYTES = 25 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 200
MAX_ZIP_MEMBERS = 2_000
MAX_PDF_PAGES = 500
MAX_PDF_PAGE_CHARS = 500_000
MAX_WORKSHEETS = 200
MAX_WORKSHEET_ROWS = 100_000
MAX_WORKSHEET_COLUMNS = 1_000
MAX_WORKSHEET_CELLS = 1_000_000
MAX_WORKBOOK_CELLS = 2_000_000
MAX_DOCX_BLOCKS = 100_000
MAX_DOCX_TABLE_ROWS = 10_000
MAX_DOCX_TABLE_COLUMNS = 256
MAX_DOCX_TABLE_CELLS = 100_000
MAX_EXTRACTED_CHARS = 5_000_000
MAX_CHUNK_CHARS = 1_000_000
MAX_CHUNKS = 10_000
MAX_SUBPROCESS_OUTPUT_BYTES = 20 * 1024 * 1024
MAX_ANTIWORD_STDERR_BYTES = 64 * 1024


class ProjectInitFileParseError(Exception):
    """Raised when a supported project-init file cannot be parsed."""


class UnsupportedProjectInitFile(Exception):
    """Raised when a project-init file has an unsupported extension."""


@dataclass(frozen=True)
class SourceChunk:
    file_name: str
    location: str
    text: str

    @property
    def source_label(self) -> str:
        return f"{self.file_name} · {self.location}"


PdfReaderFactory = Callable[[Path], Any]


def parse_project_init_file(
    path: Path,
    original_name: str,
    *,
    pdf_reader_factory: PdfReaderFactory | None = None,
) -> list[SourceChunk]:
    """Parse a supported file and normalize parser failures for callers."""
    path = Path(path)
    extension = Path(original_name).suffix.lower()
    if extension not in {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".txt"}:
        raise UnsupportedProjectInitFile(f"不支持的项目初始化文件：{original_name}")

    try:
        _validate_input_size(path, original_name)
        if extension in {".docx", ".xlsx"}:
            _validate_ooxml_archive(path, original_name)
        if extension == ".pdf":
            chunks = _parse_pdf(
                path,
                original_name,
                pdf_reader_factory or PdfReader,
            )
        elif extension == ".doc":
            chunks = _parse_doc(path, original_name)
        elif extension == ".docx":
            chunks = _parse_docx(path, original_name)
        elif extension == ".xls":
            chunks = _parse_xls(path, original_name)
        elif extension == ".xlsx":
            chunks = _parse_xlsx(path, original_name)
        else:
            chunks = _parse_txt(path, original_name)
        return _enforce_chunk_limits(chunks, original_name)
    except (ProjectInitFileParseError, UnsupportedProjectInitFile):
        raise
    except Exception as exc:
        raise ProjectInitFileParseError(
            f"无法解析项目初始化文件：{original_name}"
        ) from exc


def _parse_pdf(
    path: Path,
    original_name: str,
    reader_factory: PdfReaderFactory,
) -> Iterable[SourceChunk]:
    reader = reader_factory(path)
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ProjectInitFileParseError(f"PDF 页数超过限制：{original_name}")
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if len(text) > MAX_PDF_PAGE_CHARS:
            raise ProjectInitFileParseError(f"PDF 单页文本超过限制：{original_name}")
        yield SourceChunk(original_name, f"第 {page_number} 页", text)


def _parse_doc(path: Path, original_name: str) -> Iterable[SourceChunk]:
    resolved_path = path.resolve()
    process = subprocess.Popen(
        ["antiword", "-w", "0", str(resolved_path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        shell=False,
    )
    stdout = bytearray()
    stderr = bytearray()
    output_exceeded = threading.Event()
    readers = [
        threading.Thread(
            target=_read_subprocess_stream,
            args=(
                process.stdout,
                MAX_SUBPROCESS_OUTPUT_BYTES,
                stdout,
                output_exceeded,
                process,
            ),
            daemon=True,
        ),
        threading.Thread(
            target=_read_subprocess_stream,
            args=(
                process.stderr,
                MAX_ANTIWORD_STDERR_BYTES,
                stderr,
                output_exceeded,
                process,
            ),
            daemon=True,
        ),
    ]
    for reader in readers:
        reader.start()
    try:
        try:
            returncode = process.wait(timeout=30)
        except subprocess.TimeoutExpired as exc:
            _kill_process(process)
            process.wait()
            raise ProjectInitFileParseError(
                f"antiword 执行超时：{original_name}"
            ) from exc
    finally:
        if process.poll() is None:
            _kill_process(process)
        for reader in readers:
            reader.join()
        process.stdout.close()
        process.stderr.close()

    if output_exceeded.is_set():
        raise ProjectInitFileParseError(f"antiword 输出超过限制：{original_name}")
    if returncode != 0:
        raise ProjectInitFileParseError(
            f"无法解析项目初始化文件：{original_name}"
        )
    return _chunk_lines(_decode_antiword_output(bytes(stdout), original_name), original_name)


def _read_subprocess_stream(
    stream: Any,
    byte_limit: int,
    output: bytearray,
    output_exceeded: threading.Event,
    process: Any,
) -> None:
    bytes_read = 0
    while True:
        chunk = stream.read(max(1, min(64 * 1024, byte_limit - bytes_read + 1)))
        if not chunk:
            return
        bytes_read += len(chunk)
        if bytes_read > byte_limit:
            output_exceeded.set()
            _kill_process(process)
            return
        output.extend(chunk)


def _kill_process(process: Any) -> None:
    try:
        process.kill()
    except OSError:
        pass


def _decode_antiword_output(output: bytes, original_name: str) -> str:
    try:
        return output.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return output.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise ProjectInitFileParseError(
                f"antiword 输出编码无效：{original_name}"
            ) from exc


class _IncrementalChunkLimits:
    def __init__(self, original_name: str) -> None:
        self.original_name = original_name
        self.extracted_chars = 0
        self.chunk_count = 0

    def builder(self) -> "_IncrementalChunkBuilder":
        return _IncrementalChunkBuilder(self)

    def validate_non_empty_fragment(self, text: str, worksheet_chars: int) -> int:
        if self.chunk_count >= MAX_CHUNKS:
            raise ProjectInitFileParseError(
                f"文本块数量超过限制：{self.original_name}"
            )
        rendered_length = len(text)
        if rendered_length > MAX_CHUNK_CHARS:
            raise ProjectInitFileParseError(
                f"单个文本块超过限制：{self.original_name}"
            )
        worksheet_chars += rendered_length
        if self.extracted_chars + worksheet_chars > MAX_EXTRACTED_CHARS:
            raise ProjectInitFileParseError(
                f"提取文本超过限制：{self.original_name}"
            )
        return worksheet_chars


class _IncrementalChunkBuilder:
    def __init__(self, limits: _IncrementalChunkLimits) -> None:
        self.limits = limits
        self.buffer = io.StringIO()
        self.length = 0
        self.has_content = False

    def append(self, text: str) -> None:
        has_content = self.has_content or bool(text.strip())
        next_length = self.length + len(text)
        if has_content and next_length > MAX_CHUNK_CHARS:
            raise ProjectInitFileParseError(
                f"单个文本块超过限制：{self.limits.original_name}"
            )
        if (
            has_content
            and self.limits.extracted_chars + next_length > MAX_EXTRACTED_CHARS
        ):
            raise ProjectInitFileParseError(
                f"提取文本超过限制：{self.limits.original_name}"
            )
        if text.strip() and self.limits.chunk_count >= MAX_CHUNKS:
            raise ProjectInitFileParseError(
                f"文本块数量超过限制：{self.limits.original_name}"
            )
        self.buffer.write(text)
        self.length = next_length
        self.has_content = has_content

    def finish(self, location: str) -> SourceChunk | None:
        if not self.has_content:
            return None
        if self.limits.chunk_count >= MAX_CHUNKS:
            raise ProjectInitFileParseError(
                f"文本块数量超过限制：{self.limits.original_name}"
            )
        self.limits.extracted_chars += self.length
        self.limits.chunk_count += 1
        return SourceChunk(
            self.limits.original_name,
            location,
            self.buffer.getvalue(),
        )


def _parse_docx(path: Path, original_name: str) -> Iterable[SourceChunk]:
    document = Document(path)
    limits = _IncrementalChunkLimits(original_name)
    paragraph_batch_count = 0
    paragraph_count = 0
    table_count = 0
    block_count = 0
    paragraph_builder = limits.builder()

    def flush_paragraphs() -> SourceChunk | None:
        nonlocal paragraph_batch_count, paragraph_builder
        if not paragraph_batch_count:
            return None
        start = paragraph_count - paragraph_batch_count + 1
        chunk = paragraph_builder.finish(f"第 {start}-{paragraph_count} 段")
        paragraph_batch_count = 0
        paragraph_builder = limits.builder()
        return chunk

    for block in document.iter_inner_content():
        block_count += 1
        if block_count > MAX_DOCX_BLOCKS:
            raise ProjectInitFileParseError(f"DOCX 内容块超过限制：{original_name}")
        if isinstance(block, Paragraph):
            paragraph_count += 1
            if paragraph_batch_count:
                paragraph_builder.append("\n")
            paragraph_builder.append(block.text)
            paragraph_batch_count += 1
            if paragraph_batch_count == 20:
                chunk = flush_paragraphs()
                if chunk is not None:
                    yield chunk
        elif isinstance(block, Table):
            chunk = flush_paragraphs()
            if chunk is not None:
                yield chunk
            table_count += 1
            chunk = _docx_table_chunk(block, table_count, original_name, limits)
            if chunk is not None:
                yield chunk
    chunk = flush_paragraphs()
    if chunk is not None:
        yield chunk


def _docx_table_chunk(
    table: Table,
    table_number: int,
    original_name: str,
    limits: _IncrementalChunkLimits,
) -> SourceChunk | None:
    row_count = len(table.rows)
    column_count = max((len(row.cells) for row in table.rows), default=0)
    cell_count = sum(len(row.cells) for row in table.rows)
    if row_count > MAX_DOCX_TABLE_ROWS:
        raise ProjectInitFileParseError(f"DOCX 表格行数超过限制：{original_name}")
    if column_count > MAX_DOCX_TABLE_COLUMNS:
        raise ProjectInitFileParseError(f"DOCX 表格列数超过限制：{original_name}")
    if cell_count > MAX_DOCX_TABLE_CELLS:
        raise ProjectInitFileParseError(f"DOCX 表格单元格超过限制：{original_name}")

    end_cell = f"{get_column_letter(max(column_count, 1))}{max(row_count, 1)}"
    builder = limits.builder()
    for row_index, row in enumerate(table.rows):
        if row_index:
            builder.append("\n")
        for column_index, cell in enumerate(row.cells):
            if column_index:
                builder.append("\t")
            builder.append(cell.text)
    return builder.finish(f"第 {table_number} 个表格 A1:{end_cell}")


def _parse_xls(path: Path, original_name: str) -> Iterable[SourceChunk]:
    workbook = xlrd.open_workbook(str(path))
    try:
        sheets = workbook.sheets()
        if len(sheets) > MAX_WORKSHEETS:
            raise ProjectInitFileParseError(f"工作表数量超过限制：{original_name}")
        limits = _IncrementalChunkLimits(original_name)
        workbook_cells = 0
        for sheet in sheets:
            workbook_cells += _validate_worksheet_dimensions(
                sheet.nrows,
                sheet.ncols,
                original_name,
            )
            if workbook_cells > MAX_WORKBOOK_CELLS:
                raise ProjectInitFileParseError(f"工作簿单元格超过限制：{original_name}")
            def rows() -> Iterable[Sequence[Any]]:
                for row_index in range(sheet.nrows):
                    yield [
                        _xls_cell_value(
                            sheet.cell(row_index, column_index),
                            workbook.datemode,
                        )
                        for column_index in range(sheet.ncols)
                    ]

            chunk = _worksheet_chunk(sheet.name, rows, original_name, limits)
            if chunk is not None:
                yield chunk
    finally:
        release_resources = getattr(workbook, "release_resources", None)
        if callable(release_resources):
            release_resources()


def _parse_xlsx(path: Path, original_name: str) -> Iterable[SourceChunk]:
    cached_workbook = load_workbook(path, read_only=True, data_only=True)
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(formula_workbook.worksheets) > MAX_WORKSHEETS:
            raise ProjectInitFileParseError(f"工作表数量超过限制：{original_name}")
        if len(cached_workbook.worksheets) != len(formula_workbook.worksheets):
            raise ProjectInitFileParseError(f"工作簿结构不一致：{original_name}")
        limits = _IncrementalChunkLimits(original_name)
        workbook_cells = 0
        for cached_sheet, formula_sheet in zip(
            cached_workbook.worksheets,
            formula_workbook.worksheets,
        ):
            max_row = max(cached_sheet.max_row, formula_sheet.max_row)
            max_column = max(cached_sheet.max_column, formula_sheet.max_column)
            workbook_cells += _validate_worksheet_dimensions(
                max_row,
                max_column,
                original_name,
            )
            if workbook_cells > MAX_WORKBOOK_CELLS:
                raise ProjectInitFileParseError(f"工作簿单元格超过限制：{original_name}")
            def rows() -> Iterable[Sequence[Any]]:
                cached_rows = cached_sheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                )
                formula_rows = formula_sheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                )
                for cached_row, formula_row in zip(cached_rows, formula_rows):
                    yield [
                        formula_cell.value
                        if formula_cell.data_type == "f"
                        else cached_cell.value
                        if cached_cell.value is not None
                        else formula_cell.value
                        for cached_cell, formula_cell in zip(cached_row, formula_row)
                    ]

            chunk = _worksheet_chunk(formula_sheet.title, rows, original_name, limits)
            if chunk is not None:
                yield chunk
    finally:
        cached_workbook.close()
        formula_workbook.close()


def _validate_worksheet_dimensions(
    row_count: int,
    column_count: int,
    original_name: str,
) -> int:
    if row_count > MAX_WORKSHEET_ROWS:
        raise ProjectInitFileParseError(f"工作表行数超过限制：{original_name}")
    if column_count > MAX_WORKSHEET_COLUMNS:
        raise ProjectInitFileParseError(f"工作表列数超过限制：{original_name}")
    cell_count = row_count * column_count
    if cell_count > MAX_WORKSHEET_CELLS:
        raise ProjectInitFileParseError(f"工作表单元格超过限制：{original_name}")
    return cell_count


def _xls_cell_value(cell: Any, datemode: int) -> Any:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, datemode)
    return cell.value


def _parse_txt(path: Path, original_name: str) -> Iterable[SourceChunk]:
    with path.open("rb") as source:
        raw = source.read(MAX_INPUT_BYTES + 1)
    if len(raw) > MAX_INPUT_BYTES:
        raise ProjectInitFileParseError(f"文件大小超过限制：{original_name}")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
    return _chunk_lines(text, original_name)


def _validate_input_size(path: Path, original_name: str) -> None:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ProjectInitFileParseError(f"文件大小超过限制：{original_name}")


def _validate_ooxml_archive(path: Path, original_name: str) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise ProjectInitFileParseError(f"压缩包成员数超过限制：{original_name}")
        total_uncompressed = 0
        for member in members:
            total_uncompressed += member.file_size
            if total_uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ProjectInitFileParseError(f"压缩包解压大小超过限制：{original_name}")
            if member.file_size:
                if member.compress_size == 0:
                    raise ProjectInitFileParseError(f"压缩比超过限制：{original_name}")
                if member.file_size / member.compress_size > MAX_ZIP_COMPRESSION_RATIO:
                    raise ProjectInitFileParseError(f"压缩比超过限制：{original_name}")


def _enforce_chunk_limits(
    chunks: Iterable[SourceChunk],
    original_name: str,
) -> list[SourceChunk]:
    bounded = []
    extracted_chars = 0
    for chunk in chunks:
        if not chunk.text.strip():
            continue
        if len(chunk.text) > MAX_CHUNK_CHARS:
            raise ProjectInitFileParseError(f"单个文本块超过限制：{original_name}")
        extracted_chars += len(chunk.text)
        if extracted_chars > MAX_EXTRACTED_CHARS:
            raise ProjectInitFileParseError(f"提取文本超过限制：{original_name}")
        if len(bounded) >= MAX_CHUNKS:
            raise ProjectInitFileParseError(f"文本块数量超过限制：{original_name}")
        bounded.append(chunk)
    return bounded


def _chunk_lines(text: str, original_name: str) -> Iterable[SourceChunk]:
    return _chunk_numbered_text(_iter_lines(text), 80, "行", original_name)


def _iter_lines(text: str) -> Iterable[str]:
    source = io.StringIO(text, newline=None)
    try:
        for line in source:
            yield line.rstrip("\r\n")
    finally:
        source.close()


def _chunk_numbered_text(
    values: Iterable[str],
    chunk_size: int,
    unit: str,
    original_name: str,
) -> Iterable[SourceChunk]:
    batch = []
    start_index = 1
    for value in values:
        batch.append(value)
        if len(batch) == chunk_size:
            end_index = start_index + len(batch) - 1
            yield SourceChunk(
                original_name,
                f"第 {start_index}-{end_index} {unit}",
                "\n".join(batch),
            )
            start_index = end_index + 1
            batch.clear()
    if batch:
        end_index = start_index + len(batch) - 1
        yield SourceChunk(
            original_name,
            f"第 {start_index}-{end_index} {unit}",
            "\n".join(batch),
        )


def _worksheet_chunk(
    sheet_name: str,
    rows: Callable[[], Iterable[Sequence[Any]]] | Iterable[Sequence[Any]],
    original_name: str,
    limits: _IncrementalChunkLimits | None = None,
) -> SourceChunk | None:
    row_factory = rows if callable(rows) else lambda: rows
    limits = limits or _IncrementalChunkLimits(original_name)
    min_row = max_row = min_column = max_column = None
    worksheet_chars = 0
    for row_index, row in enumerate(row_factory()):
        for column_index, value in enumerate(row):
            if not _is_non_empty_cell(value):
                continue
            worksheet_chars = limits.validate_non_empty_fragment(
                _render_cell(value),
                worksheet_chars,
            )
            min_row = row_index if min_row is None else min(min_row, row_index)
            max_row = row_index if max_row is None else max(max_row, row_index)
            min_column = (
                column_index
                if min_column is None
                else min(min_column, column_index)
            )
            max_column = (
                column_index
                if max_column is None
                else max(max_column, column_index)
            )
    if min_row is None or max_row is None or min_column is None or max_column is None:
        return None

    builder = limits.builder()
    rendered_row_count = 0
    for row_index, row in enumerate(row_factory()):
        if row_index < min_row or row_index > max_row:
            continue
        if rendered_row_count:
            builder.append("\n")
        for column_index in range(min_column, max_column + 1):
            if column_index > min_column:
                builder.append("\t")
            value = row[column_index] if column_index < len(row) else None
            builder.append(_render_cell(value))
        rendered_row_count += 1
    location = (
        f"{quote_sheetname(sheet_name)}!"
        f"{get_column_letter(min_column + 1)}{min_row + 1}:"
        f"{get_column_letter(max_column + 1)}{max_row + 1}"
    )
    return builder.finish(location)


def _is_non_empty_cell(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _render_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)
