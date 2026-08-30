"""Conservative structural profiling for Excel project-init attachments."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def profile_project_init_workbook(path: Path, original_name: str) -> dict[str, Any]:
    """Return safe routing metadata without exposing workbook cell values.

    This is deliberately a structural profile: it records layout traits that
    make text-only extraction lossy, but never copies business content into
    logs, snapshots, or routing metadata.
    """

    if Path(original_name).suffix.lower() != ".xlsx":
        return {
            "risk_level": "low",
            "signals": [],
            "summary": {"inspection": "not_applicable"},
        }

    try:
        workbook = load_workbook(
            io.BytesIO(Path(path).read_bytes()),
            read_only=False,
            data_only=False,
        )
    except Exception:
        return {
            "risk_level": "high",
            "signals": ["workbook_structure_unavailable"],
            "summary": {"inspection": "unavailable"},
        }

    try:
        hidden_sheet_count = sum(
            sheet.sheet_state != "visible" for sheet in workbook.worksheets
        )
        merged_range_count = sum(
            len(sheet.merged_cells.ranges) for sheet in workbook.worksheets
        )
        formula_cell_count = sum(
            1
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        )
        summary = {
            "worksheet_count": len(workbook.worksheets),
            "visible_sheet_count": len(workbook.worksheets) - hidden_sheet_count,
            "hidden_sheet_count": hidden_sheet_count,
            "merged_range_count": merged_range_count,
            "formula_cell_count": formula_cell_count,
        }
    finally:
        workbook.close()

    signals: list[str] = []
    if hidden_sheet_count:
        signals.append("hidden_sheets")
    if merged_range_count:
        signals.append("merged_cells")
    if formula_cell_count:
        signals.append("formulas")

    if hidden_sheet_count or len(signals) >= 2:
        risk_level = "high"
    elif signals:
        risk_level = "medium"
    else:
        risk_level = "low"
    return {"risk_level": risk_level, "signals": signals, "summary": summary}
