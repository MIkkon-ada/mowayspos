"""Render bounded, visible Excel worksheet views for vision analysis."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


MAX_RENDERED_CELLS = 50_000
MAX_ROWS_PER_IMAGE = 50
MAX_COLUMNS_PER_IMAGE = 24
CELL_WIDTH = 172
CELL_HEIGHT = 46
_BORDER = "#CBD5E1"
_HEADER_BACKGROUND = "#E0F2FE"
_BODY_BACKGROUND = "#FFFFFF"
_TEXT = "#0F172A"


class WorkbookRenderLimitError(ValueError):
    """Raised when a workbook cannot be rendered within bounded vision limits."""


def render_workbook_images(
    path: Path,
    output_directory: Path,
    *,
    max_images: int = 8,
) -> list[Path]:
    """Render each visible worksheet range into readable PNG pages.

    Callers own ``output_directory`` and must remove it after model invocation.
    The renderer never persists cell content outside those temporary PNG files.
    """

    if max_images < 1:
        raise ValueError("max_images must be positive")
    workbook = load_workbook(
        io.BytesIO(Path(path).read_bytes()),
        read_only=False,
        data_only=False,
    )
    try:
        pages: list[tuple[object, tuple[int, int, int, int]]] = []
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            bounds = _used_bounds(sheet)
            if bounds is None:
                continue
            min_row, max_row, min_column, max_column = bounds
            cell_count = (max_row - min_row + 1) * (max_column - min_column + 1)
            if cell_count > MAX_RENDERED_CELLS:
                raise WorkbookRenderLimitError("workbook render cell limit exceeded")
            for row_start in range(min_row, max_row + 1, MAX_ROWS_PER_IMAGE):
                row_end = min(max_row, row_start + MAX_ROWS_PER_IMAGE - 1)
                for column_start in range(min_column, max_column + 1, MAX_COLUMNS_PER_IMAGE):
                    column_end = min(max_column, column_start + MAX_COLUMNS_PER_IMAGE - 1)
                    pages.append((sheet, (row_start, row_end, column_start, column_end)))
        if len(pages) > max_images:
            raise WorkbookRenderLimitError("workbook render image count limit exceeded")

        output = Path(output_directory)
        output.mkdir(parents=True, exist_ok=True)
        return [
            _render_page(sheet, bounds, output / f"workbook-{index:02d}.png")
            for index, (sheet, bounds) in enumerate(pages, start=1)
        ]
    finally:
        workbook.close()


def _used_bounds(sheet) -> tuple[int, int, int, int] | None:
    min_row = min_column = None
    max_row = max_column = None
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or not str(cell.value).strip():
                continue
            min_row = cell.row if min_row is None else min(min_row, cell.row)
            max_row = cell.row if max_row is None else max(max_row, cell.row)
            min_column = cell.column if min_column is None else min(min_column, cell.column)
            max_column = cell.column if max_column is None else max(max_column, cell.column)
    if min_row is None or max_row is None or min_column is None or max_column is None:
        return None
    return min_row, max_row, min_column, max_column


def _render_page(sheet, bounds: tuple[int, int, int, int], destination: Path) -> Path:
    min_row, max_row, min_column, max_column = bounds
    row_count = max_row - min_row + 1
    column_count = max_column - min_column + 1
    image = Image.new(
        "RGB",
        (column_count * CELL_WIDTH + 1, row_count * CELL_HEIGHT + 1),
        "white",
    )
    drawing = ImageDraw.Draw(image)
    font = _font()
    merged_anchors = {
        (cell_range.min_row, cell_range.min_col): cell_range
        for cell_range in sheet.merged_cells.ranges
    }
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            x0 = (column - min_column) * CELL_WIDTH
            y0 = (row - min_row) * CELL_HEIGHT
            x1 = x0 + CELL_WIDTH
            y1 = y0 + CELL_HEIGHT
            drawing.rectangle(
                (x0, y0, x1, y1),
                fill=_HEADER_BACKGROUND if row == min_row else _BODY_BACKGROUND,
                outline=_BORDER,
            )
            merged = merged_anchors.get((row, column))
            if merged is not None:
                x1 = min(
                    (merged.max_col - min_column + 1) * CELL_WIDTH,
                    image.width - 1,
                )
                y1 = min(
                    (merged.max_row - min_row + 1) * CELL_HEIGHT,
                    image.height - 1,
                )
                drawing.rectangle((x0, y0, x1, y1), fill=_BODY_BACKGROUND, outline=_BORDER)
            value = sheet.cell(row, column).value
            if value is not None and str(value).strip() and (merged is None or merged.min_row == row and merged.min_col == column):
                _draw_cell_text(drawing, str(value), (x0, y0, x1, y1), font)
    image.save(destination, format="PNG", optimize=True)
    return destination


def _draw_cell_text(drawing: ImageDraw.ImageDraw, value: str, box: tuple[int, int, int, int], font) -> None:
    x0, y0, x1, y1 = box
    text = " ".join(value.replace("\r", " ").replace("\n", " ").split())[:120]
    available_width = max(1, x1 - x0 - 12)
    while text and drawing.textbbox((0, 0), text, font=font)[2] > available_width:
        text = text[:-1]
    if text:
        drawing.text((x0 + 6, y0 + 12), text, fill=_TEXT, font=font)


def _font():
    for candidate in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
    ):
        try:
            return ImageFont.truetype(candidate, 18)
        except OSError:
            continue
    return ImageFont.load_default()
