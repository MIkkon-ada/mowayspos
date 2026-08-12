"""Parse the company's standard meeting-minutes Word layout without AI inference."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document


_DATE_PATTERN = re.compile(r"(20\d{2}-\d{2}-\d{2})")
_AGENDA_NUMBER = re.compile(r"^\s*\d+[、.．]\s*")
_TITLE_SUFFIX = re.compile(r"\s*[会议纪要]+\s*$")


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _key(value: str) -> str:
    return re.sub(r"\s+", "", _clean(value))


def _table_as_records(table) -> list[dict[str, str]]:
    rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
    if len(rows) < 2 or not all(rows[0]):
        return []
    headers = rows[0]
    return [
        {header: values[index] if index < len(values) else "" for index, header in enumerate(headers)}
        for values in rows[1:]
        if any(values)
    ]


def _table_values(document) -> dict[str, str]:
    values: dict[str, str] = {}
    for table in document.tables:
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if len(cells) not in {2, 4}:
                continue
            for index in range(0, len(cells), 2):
                if cells[index] and cells[index + 1]:
                    values[_key(cells[index])] = cells[index + 1]
    return values


def _find_table(document, required_headers: set[str]) -> list[dict[str, str]]:
    for table in document.tables:
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        if rows and required_headers.issubset(set(rows[0])):
            return _table_as_records(table)
    return []


def _section_lines(paragraphs: list[str], heading: str, next_heading: str) -> list[str]:
    try:
        start = next(index for index, line in enumerate(paragraphs) if line.startswith(heading)) + 1
    except StopIteration:
        return []
    lines: list[str] = []
    for line in paragraphs[start:]:
        if line.startswith(next_heading):
            break
        if line:
            lines.append(line)
    return lines


def _first_value(values: dict[str, str], *labels: str) -> str:
    for label in labels:
        if values.get(_key(label)):
            return values[_key(label)]
    return ""


def parse_standard_meeting_minutes(filename: str, content: bytes) -> dict[str, object]:
    """Return a fidelity-first structure when a Word file matches the standard layout."""
    if Path(filename or "").suffix.lower() != ".docx":
        return {"is_standard_minutes": False}
    try:
        document = Document(BytesIO(content))
    except Exception:
        return {"is_standard_minutes": False}

    paragraphs = [_clean(paragraph.text) for paragraph in document.paragraphs]
    paragraphs = [line for line in paragraphs if line]
    values = _table_values(document)
    agenda_lines = _section_lines(paragraphs, "一、会议议程", "二、会议小结与决议")
    summary_lines = _section_lines(paragraphs, "二、会议小结与决议", "三、待办事项跟踪")
    current_action_items = _find_table(document, {"编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"})
    prior_action_items = _find_table(document, {"编号", "上周事项", "负责人", "状态", "本周进展/说明"})
    footer = " ".join(paragraphs[-4:])
    copied_to_match = re.search(r"抄送\s*[:：]?\s*(.+)$", footer)
    organizer_match = re.search(r"整理人\s*[:：]?\s*([^抄送]+)", footer)
    required_basic_fields = (
        _first_value(values, "会议时间"),
        _first_value(values, "会议地点"),
        _first_value(values, "会议类型"),
        _first_value(values, "会议主持人", "主持人"),
        _first_value(values, "与会者", "参会人员"),
    )
    organizer = _first_value(values, "整理人") or (organizer_match.group(1).strip() if organizer_match else "")
    copied_to = _first_value(values, "抄送") or (copied_to_match.group(1).strip() if copied_to_match else "")
    is_standard_minutes = bool(
        agenda_lines
        and summary_lines
        and current_action_items
        and prior_action_items
        and all(required_basic_fields)
        and organizer
        and copied_to
    )
    if not is_standard_minutes:
        return {"is_standard_minutes": False}

    title = paragraphs[0] if paragraphs else ""
    title = _TITLE_SUFFIX.sub("", title).strip()
    date_source = _first_value(values, "会议时间")
    date_match = _DATE_PATTERN.search(date_source) or _DATE_PATTERN.search(" ".join(paragraphs[:3]))
    return {
        "is_standard_minutes": True,
        "title": title,
        "meeting_date": date_match.group(1) if date_match else "",
        "location": _first_value(values, "会议地点"),
        "meeting_type": _first_value(values, "会议类型"),
        "host": _first_value(values, "会议主持人", "主持人"),
        "participants": _first_value(values, "与会者", "参会人员"),
        "organizer": organizer,
        "copied_to": copied_to,
        "agenda_items": [_AGENDA_NUMBER.sub("", line) for line in agenda_lines],
        "summary": "\n".join(summary_lines),
        "current_action_items": current_action_items,
        "prior_action_items": prior_action_items,
    }
