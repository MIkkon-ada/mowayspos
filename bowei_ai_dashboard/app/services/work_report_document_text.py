"""Bounded text extraction for work-report source documents.

This module intentionally accepts bytes and returns plain text only.  It does
not write uploaded content to disk or make any business/AI decisions.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 5_000
MAX_ARCHIVE_MEMBERS = 2_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_PDF_PAGES = 100
ALLOWED_SUFFIXES = {".docx", ".pdf", ".xlsx", ".pptx"}


class WorkReportDocumentTextError(ValueError):
    """Raised when a document cannot safely provide work-report text."""


def extract_work_report_document_text(filename: str, content: bytes) -> str:
    """Extract bounded readable text from one supported work-report document."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise WorkReportDocumentTextError("仅支持 Word、PDF、Excel、PPT 文档")
    if not content:
        raise WorkReportDocumentTextError("文档内容为空")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise WorkReportDocumentTextError("文档不能超过 20 MB")

    try:
        if suffix == ".docx":
            _validate_office_archive(content)
            text = _extract_docx(content)
        elif suffix == ".xlsx":
            _validate_office_archive(content)
            text = _extract_xlsx(content)
        elif suffix == ".pptx":
            _validate_office_archive(content)
            text = _extract_pptx(content)
        else:
            text = _extract_pdf(content)
    except WorkReportDocumentTextError:
        raise
    except Exception as exc:
        raise WorkReportDocumentTextError("文档无法解析，请确认文件未损坏且未加密") from exc

    normalized = _normalize_text(text)
    if not normalized:
        raise WorkReportDocumentTextError("文档中没有可解析的文字内容")
    if len(normalized) > MAX_EXTRACTED_CHARS:
        raise WorkReportDocumentTextError("文档解析内容不能超过 5000 字，请上传相关节选")
    return normalized


def _validate_office_archive(content: bytes) -> None:
    """Reject archive structures that are unsuitable for in-memory parsing."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ARCHIVE_MEMBERS:
                raise WorkReportDocumentTextError("文档无法解析，请确认文件未损坏且未加密")
            if any(info.file_size > MAX_DOCUMENT_BYTES for info in infos):
                raise WorkReportDocumentTextError("文档内容过大")
            if sum(info.file_size for info in infos) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise WorkReportDocumentTextError("文档内容过大")
    except BadZipFile as exc:
        raise WorkReportDocumentTextError("文档无法解析，请确认文件未损坏且未加密") from exc


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                lines.append("\t".join(values))
    return "\n".join(lines)


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise WorkReportDocumentTextError("文档无法解析，请确认文件未损坏且未加密")
    return _extract_bounded_pdf_text(reader)


def _extract_bounded_pdf_text(reader: PdfReader) -> str:
    if len(reader.pages) > MAX_PDF_PAGES:
        raise WorkReportDocumentTextError("PDF 页数不能超过 100 页")

    lines: list[str] = []
    extracted_chars = 0
    for page in reader.pages:
        page_fragments: list[str] = []

        def collect_text(fragment: str, *_: object) -> None:
            nonlocal extracted_chars
            extracted_chars += len(fragment)
            if extracted_chars > MAX_EXTRACTED_CHARS:
                raise WorkReportDocumentTextError("文档解析内容不能超过 5000 字，请上传相关章节")
            page_fragments.append(fragment)

        page.extract_text(visitor_text=collect_text)
        lines.append("".join(page_fragments))
    return "\n".join(lines)


def _extract_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            sheet_lines: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    sheet_lines.append("\t".join(values))
            if sheet_lines:
                lines.append(f"[{worksheet.title}]")
                lines.extend(sheet_lines)
        return "\n".join(lines)
    finally:
        workbook.close()


def _extract_pptx(content: bytes) -> str:
    deck = Presentation(BytesIO(content))
    lines: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    slide_lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        slide_lines.append("\t".join(values))
        if slide_lines:
            lines.append(f"[第 {index} 页]")
            lines.extend(slide_lines)
    return "\n".join(lines)


def _normalize_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
