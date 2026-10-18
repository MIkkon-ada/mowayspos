from io import BytesIO

from docx import Document

from app import models
from app.services.meeting_minutes_export import build_meeting_minutes_docx


def test_export_uses_formal_minutes_sections_and_action_tracking_table():
    meeting = models.Meeting(
        title="AI升级项目周会会议纪要",
        related_special_project="AI升级计划",
        meeting_date="2026-07-27",
        location="线上会议",
        meeting_type="周会",
        host="杨宇帆",
        participants="吴肖、刘万超",
        organizer="吴肖",
        copied_to="AI升级计划项目组",
        agenda_items_json='["回顾本周进展", "确认下周行动项"]',
        summary="本周工作按计划推进。",
        decision_items_json='[{"title":"项目管理平台","bullets":["确认本周先完成计划分解"]}]',
        task_list_json='[{"编号":"本周-01","会议安排事项":"完成计划分解","负责人":"温会林","追踪人":"杨宇帆","完成时限":"2026-08-03","来源/备注":"会议决议"}]',
        prior_action_items_json='[{"content":"回顾上周进展","owner":"吴肖","status":"已完成","progress_note":"已形成纪要"}]',
        risk_items_json='[{"content":"接口依赖待确认"}]',
    )

    document = Document(BytesIO(build_meeting_minutes_docx(
        meeting,
        source_name="原始会议纪要.docx",
        completed_items=[{"content": "本周已完成计划核对"}],
    )))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]

    assert "一、会议议程" in paragraphs
    assert "二、会议讨论与结论" in paragraphs
    assert "三、会议行动项" in paragraphs
    assert "（一）本周待办事项" in paragraphs
    assert "（二）上周待办追踪" in paragraphs
    assert "本周已确认进展" in paragraphs
    assert "风险与待确认" in paragraphs
    assert "四、风险与待确认" not in paragraphs
    assert paragraphs.index("二、会议讨论与结论") < paragraphs.index("本周已确认进展") < paragraphs.index("三、会议行动项")
    assert ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"] in table_text
    assert ["本周-01", "完成计划分解", "温会林", "杨宇帆", "2026-08-03", "会议决议"] in table_text
    assert ["追踪-01", "回顾上周进展", "吴肖", "已完成", "已形成纪要"] in table_text
    assert ["整理人", "吴肖", "抄送", "AI升级计划项目组"] in table_text
