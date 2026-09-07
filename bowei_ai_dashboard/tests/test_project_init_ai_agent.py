import json

import pytest
from openpyxl import Workbook
from pydantic import ValidationError
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.ai.contracts import AIInvocationContext, AIUpstreamError, Capability
from app.ai.repository import AIConfigurationRepository
from app.ai.service import AIService
from app.database import Base
from app.services.project_init_ai_agent import (
    AgentTask,
    Evidence,
    ProjectInitAiEmptyResult,
    ProjectInitAiError,
    ProjectInitAiInvalidDraft,
    _merge_tasks,
    _context_prompt,
    _final_merge_prompt,
    _parse_json_response,
    generate_project_init_draft,
)
from app.services.project_init_file_parser import SourceChunk, parse_project_init_file


TEST_FERNET_KEY = "m6F5dBXMRy1ZOQ4Dv_rwuPhtchxZzTCBuRUg-hxeF6U="


class SequencedChatAdapters:
    def __init__(self, responses: dict[str, str]) -> None:
        self.responses = responses

    def complete_chat(self, model, _api_key, _prompt, *, timeout_seconds):
        assert timeout_seconds == 30
        return self.responses[model.code]


def chunk(text: str, *, name: str = "plan.txt", location: str = "lines 1-2") -> dict:
    return {"attachment_id": 7, "file_name": name, "location": location, "text": text}


def source_chunk(text: str, *, name: str = "plan.txt", location: str = "lines 1-2") -> SourceChunk:
    return SourceChunk(file_name=name, location=location, text=text)


def raw_task(
    *,
    title: str = "实施交付",
    owner_name: str = "张三",
    assignee_name: str = "张三",
    helper_names: list[str] | None = None,
    evidence: list[dict] | None = None,
) -> dict:
    return {
        "title": title,
        "description": "完成项目初始化工作",
        "owner_name": owner_name,
        "priority": "high",
        "status": "not_started",
        "plan_start": "2026-09-01",
        "plan_end": "2026-09-30",
        "evidence": evidence or [{"attachment_id": 7, "file_name": "plan.txt", "location": "lines 1-2", "excerpt": "实施交付"}],
        "subtasks": [
            {
                "title": "完成方案确认",
                "assignee_name": assignee_name,
                "helper_names": helper_names or [],
                "evaluation_standard": "完成确认并留痕",
                "evidence": evidence or [{"attachment_id": 7, "file_name": "plan.txt", "location": "lines 1-2", "excerpt": "实施交付"}],
            }
        ],
    }


def fake_llm(result: dict):
    def call(prompt: str, provider: str) -> str:
        assert provider == "injected"
        assert "不执行数据库、项目或成员修改" in prompt
        return json.dumps(result, ensure_ascii=False)

    return call


def test_ai_draft_accepts_an_evidence_bound_project_profile_without_work_tasks():
    source = chunk("项目名称：岗位 AI 应用优化\n项目背景：岗位知识分散\n项目目标：提升复用率")
    result = generate_project_init_draft(
        [source],
        [],
        [],
        llm_call=fake_llm({
            "project_profile": {
                "name": "岗位 AI 应用优化",
                "background": "岗位知识分散",
                "objectives": "提升复用率",
                "evidence": [{
                    "attachment_id": 7,
                    "file_name": "plan.txt",
                    "location": "lines 1-2",
                    "excerpt": "",
                }],
            },
            "tasks": [],
        }),
    )

    assert result.tasks == []
    assert result.project_profile.name == "岗位 AI 应用优化"
    assert result.project_profile.evidence[0].file_name == "plan.txt"


def test_schema_invalid_primary_response_uses_project_init_fallback_model():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    try:
        repo = AIConfigurationRepository(db, cipher_key=TEST_FERNET_KEY)
        primary = repo.create_model(
            code="primary",
            display_name="Primary",
            provider="deepseek",
            model_name="deepseek-chat",
            model_type="chat",
            base_url="https://api.example.test",
            config={},
            enabled=True,
            source="custom",
        )
        fallback = repo.create_model(
            code="fallback",
            display_name="Fallback",
            provider="dashscope",
            model_name="qwen-plus",
            model_type="chat",
            base_url="https://api.example.test",
            config={},
            enabled=True,
            source="custom",
        )
        for model in (primary, fallback):
            repo.replace_credential(model.id, api_key="test-key", app_secret=None)
        repo.save_policy(
            Capability.PROJECT_INIT_ANALYSIS,
            primary_model_id=primary.id,
            fallback_model_ids=[fallback.id],
            timeout_seconds=30,
            max_attempts=2,
            enabled=True,
        )
        invalid = raw_task()
        invalid["evidence"] = invalid["evidence"][0]
        invalid["subtasks"][0]["evidence"] = invalid["subtasks"][0]["evidence"][0]
        adapters = SequencedChatAdapters(
            {
                "primary": json.dumps({"tasks": [invalid]}, ensure_ascii=False),
                "fallback": json.dumps({"tasks": [raw_task()]}, ensure_ascii=False),
            }
        )

        result = generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            ai_service=AIService(db, adapters=adapters, cipher_key=TEST_FERNET_KEY),
            invocation_context=AIInvocationContext(resource_type="project_init", resource_id=99),
        )

        assert result.tasks[0].title == "实施交付"
        logs = db.query(models.AIInvocationLog).order_by(models.AIInvocationLog.id).all()
        assert [(log.status, log.fallback_used, log.error_code) for log in logs] == [
            ("failed", False, "AI_RESPONSE_INVALID"),
            ("succeeded", True, ""),
        ]
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_semantic_workbook_uses_deterministic_projection_when_all_models_fail():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    try:
        repo = AIConfigurationRepository(db, cipher_key=TEST_FERNET_KEY)
        primary = repo.create_model(
            code="primary",
            display_name="Primary",
            provider="deepseek",
            model_name="deepseek-chat",
            model_type="chat",
            base_url="https://api.example.test",
            config={},
            enabled=True,
            source="custom",
        )
        repo.replace_credential(primary.id, api_key="test-key", app_secret=None)
        repo.save_policy(
            Capability.PROJECT_INIT_ANALYSIS,
            primary_model_id=primary.id,
            fallback_model_ids=[],
            timeout_seconds=30,
            max_attempts=1,
            enabled=True,
        )

        class FailingAdapter:
            def complete_chat(self, _model, _api_key, _prompt, *, timeout_seconds):
                raise AIUpstreamError("AI_UPSTREAM_TIMEOUT", retryable=True)

        result = generate_project_init_draft(
            [
                chunk(
                    "主要工作\t关键任务\t目标\t推进流程\n"
                    "联合拓展\t明确客户范围\t完成客户筛选\t拜访并复盘",
                    name="工作推进表.xlsx",
                    location="Sheet1!A1:D2",
                )
            ],
            [{"id": 1, "name": "张三", "is_active": True}],
            [],
            ai_service=AIService(db, adapters=FailingAdapter(), cipher_key=TEST_FERNET_KEY),
            invocation_context=AIInvocationContext(resource_type="project_init", resource_id=99),
        )

        assert result.model_name == "structured-spreadsheet-fallback"
        assert result.tasks[0].title == "联合拓展"
        assert result.tasks[0].subtasks[0].title == "明确客户范围"
        assert result.tasks[0].goal == "完成客户筛选"
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_unique_active_person_is_bound_but_ambiguous_and_inactive_are_not():
    people = [
        {"id": 1, "name": "张三", "is_active": True},
        {"id": 2, "name": "李四", "is_active": True},
        {"id": 3, "name": "李四", "is_active": True},
        {"id": 4, "name": "王五", "is_active": False},
    ]
    result = generate_project_init_draft(
        [chunk("实施交付\n完成方案确认")],
        people,
        [],
        llm_call=fake_llm({"tasks": [raw_task(helper_names=["李四", "王五"])]}),
    )

    task = result.tasks[0]
    subtask = task.subtasks[0]
    assert task.owner_id == 1
    assert subtask.assignee_id == 1
    assert subtask.helper_ids == []
    assert {warning.code for warning in subtask.warnings} == {"ambiguous_person", "inactive_person"}


def test_full_person_snapshot_is_reduced_to_candidate_fields():
    result = generate_project_init_draft(
        [chunk("实施交付")],
        [{
            "id": 1,
            "name": "张三",
            "is_active": True,
            "department": "交付部",
            "system_role": "normal_member",
            "special_project_duty": "项目负责人",
        }],
        [],
        llm_call=fake_llm({"tasks": [raw_task()]}),
    )

    assert result.tasks[0].owner_id == 1
    assert result.tasks[0].subtasks[0].assignee_id == 1


def test_unmatched_person_name_is_preserved_without_an_id():
    result = generate_project_init_draft(
        [chunk("安排外部顾问")],
        [],
        [],
        llm_call=fake_llm({
            "tasks": [raw_task(
                owner_name="外部顾问",
                assignee_name="外部顾问",
                evidence=[{"attachment_id": 7, "file_name": "plan.txt", "location": "lines 1-2", "excerpt": "安排外部顾问"}],
            )]
        }),
    )

    task = result.tasks[0]
    assert task.owner_name == "外部顾问"
    assert task.owner_id is None
    assert task.subtasks[0].assignee_name == "外部顾问"
    assert task.subtasks[0].assignee_id is None
    assert any(warning.code == "person_not_found" for warning in task.subtasks[0].warnings)


def test_duplicate_classification_is_deterministic_and_never_merges():
    existing = [
        {"id": 10, "title": "实施交付", "subtasks": [{"id": 11, "title": "完成方案确认"}]},
        {"id": 20, "title": "完成系统上线任务", "subtasks": [{"id": 21, "title": "上线验收"}]},
    ]
    result = generate_project_init_draft(
        [chunk("实施交付\n完成系统上线")],
        [],
        existing,
        llm_call=fake_llm(
            {
                "tasks": [
                    raw_task(title="实施交付", assignee_name=""),
                    raw_task(title="完成系统上线任", assignee_name=""),
                    raw_task(title="建立风险台账", assignee_name=""),
                ]
            }
        ),
    )

    assert [task.merge_status for task in result.tasks] == [
        "definite_duplicate",
        "possible_duplicate",
        "new",
    ]
    assert result.tasks[0].duplicate_of == 10
    assert result.tasks[1].duplicate_of == 20
    assert all(task.subtasks for task in result.tasks)
    assert result.tasks[0].subtasks[0].merge_status == "definite_duplicate"
    assert result.tasks[0].subtasks[0].duplicate_of == 11


def test_source_evidence_is_preserved_and_model_text_is_not_used_as_evidence():
    evidence = [{"attachment_id": 7, "file_name": "source.docx", "location": "paragraphs 2-3", "excerpt": "原文片段"}]
    result = generate_project_init_draft(
        [chunk("原文片段", name="source.docx", location="paragraphs 2-3")],
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(evidence=evidence)]}),
    )
    item = result.tasks[0].evidence[0]
    assert item.file_name == "source.docx"
    assert item.location == "paragraphs 2-3"
    assert item.source_label == "source.docx · paragraphs 2-3"


def test_ai_draft_routes_project_profile_and_work_progress_with_traceable_evidence():
    result = generate_project_init_draft(
        [chunk(
            "项目名称：岗位 AI 应用优化\n"
            "建设背景：岗位知识分散\n"
            "项目目标：提升复用率\n"
            "预期成果：形成案例库\n"
            "项目周期：2026-07-01 至 2026-09-30\n"
            "工作模块：岗位优化\n"
            "关键任务：建立应用记录表",
            name="项目方案.xlsx",
            location="概况!A1:B8",
        )],
        [],
        [],
        llm_call=fake_llm({
            "project_profile": {
                "name": "岗位 AI 应用优化",
                "background": "岗位知识分散",
                "objectives": "提升复用率",
                "expected_outcomes": "形成案例库",
                "start_date": "2026-07-01",
                "end_date": "2026-09-30",
                "description": "",
                "confidence": 0.96,
                "evidence": [{
                    "attachment_id": 7,
                    "file_name": "项目方案.xlsx",
                    "location": "概况!A1:B8",
                    "excerpt": "项目名称：岗位 AI 应用优化",
                }],
                "warnings": [],
            },
            "tasks": [raw_task(
                title="岗位优化",
                evidence=[{
                    "attachment_id": 7,
                    "file_name": "项目方案.xlsx",
                    "location": "概况!A1:B8",
                    "excerpt": "工作模块：岗位优化",
                }],
            )],
        }),
    )

    assert result.project_profile.name == "岗位 AI 应用优化"
    assert result.project_profile.background == "岗位知识分散"
    assert result.project_profile.objectives == "提升复用率"
    assert result.project_profile.expected_outcomes == "形成案例库"
    assert result.project_profile.start_date == "2026-07-01"
    assert result.project_profile.end_date == "2026-09-30"
    assert result.project_profile.evidence[0].location == "概况!A1:B8"
    assert [task.title for task in result.tasks] == ["岗位优化"]


def test_semantic_workbook_shape_preserves_goal_acceptance_process_and_numbered_tasks(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "复制推广"
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "第三阶段五项工作的简要方案"
    sheet.append(["主要工作", "目标", "验收标准与关键成果", "关键任务", "推进流程"])
    sheet.append([
        "1. 试点岗位AI应用优化",
        "AI全面嵌入试点岗位，真实运行与优化",
        "1. 完成岗位应用\n2. 形成复盘记录",
        "1. 设定优化期间的绩效提升目标\n2. 建立应用记录表",
        "宣贯 → 试用 → 复盘 → 优化",
    ])
    workbook_path = tmp_path / "第三阶段五项工作的简要方案_v0.2.xlsx"
    workbook.save(workbook_path)
    parsed_chunks = parse_project_init_file(workbook_path, workbook_path.name)
    source = next(item for item in parsed_chunks if item.location.endswith("A3:E3"))

    data = {
        "tasks": [
            {
                "title": "1. 试点岗位AI应用优化",
                "goal": "AI全面嵌入试点岗位，真实运行与优化",
                "acceptance_criteria": "1. 完成岗位应用\n2. 形成复盘记录",
                "process": "宣贯 → 试用 → 复盘 → 优化",
                "description": "AI全面嵌入试点岗位，真实运行与优化",
                "evidence": [{"attachment_id": 7, "file_name": source.file_name, "location": source.location}],
                "subtasks": [
                    {
                        "title": "设定优化期间的绩效提升目标",
                        "evaluation_standard": "完成岗位应用",
                        "evidence": [{"attachment_id": 7, "file_name": source.file_name, "location": source.location}],
                    },
                    {
                        "title": "建立应用记录表",
                        "evaluation_standard": "形成复盘记录",
                        "evidence": [{"attachment_id": 7, "file_name": source.file_name, "location": source.location}],
                    },
                ],
            }
        ]
    }
    result = generate_project_init_draft(
        [{"attachment_id": 7, "file_name": source.file_name, "location": source.location, "text": source.text}],
        [],
        [],
        llm_call=fake_llm(data),
    )

    task = result.tasks[0]
    assert task.goal == "AI全面嵌入试点岗位，真实运行与优化"
    assert task.acceptance_criteria.startswith("1. 完成岗位应用")
    assert task.process == "宣贯 → 试用 → 复盘 → 优化"
    assert [subtask.title for subtask in task.subtasks] == [
        "设定优化期间的绩效提升目标",
        "建立应用记录表",
    ]
    assert task.evidence[0].location == source.location


def test_semantic_workbook_missing_fields_get_one_bounded_ai_repair():
    source = chunk(
        "主要工作\t目标\t验收标准与关键成果\t关键任务\t推进流程\n"
        "试点优化\t岗位真实运行\t完成验收\t1.建立记录表 2.完成复盘\t试用 → 复盘",
        name="推进表.xlsx",
        location="'复制推广'!A3:E3",
    )
    first = raw_task(
        title="试点优化",
        evidence=[{"attachment_id": 7, "file_name": "推进表.xlsx", "location": "'复制推广'!A3:E3"}],
    )
    repaired = {
        **first,
        "goal": "岗位真实运行",
        "acceptance_criteria": "完成验收",
        "process": "试用 → 复盘",
    }
    calls: list[str] = []

    def llm(prompt: str, _provider: str) -> str:
        calls.append(prompt)
        return json.dumps({"tasks": [first if len(calls) == 1 else repaired]}, ensure_ascii=False)

    result = generate_project_init_draft([source], [], [], llm_call=llm)

    assert len(calls) == 2
    assert result.tasks[0].goal == "岗位真实运行"
    assert result.tasks[0].acceptance_criteria == "完成验收"
    assert result.tasks[0].process == "试用 → 复盘"


def test_semantic_workbook_falls_back_to_source_fields_when_ai_omits_them():
    source = chunk(
        "主要工作\t目标\t验收标准与关键成果\t关键任务\t推进流程\n"
        "试点优化\t岗位真实运行\t完成验收\t1.建立记录表 2.完成复盘\t试用 → 复盘",
        name="第三阶段五项工作的简要方案_v0.2.xlsx",
        location="'复制推广'!A3:E3",
    )
    candidate = raw_task(
        title="试点优化",
        evidence=[{"attachment_id": 7, "file_name": source["file_name"], "location": source["location"]}],
    )
    calls: list[str] = []

    def llm(prompt: str, _provider: str) -> str:
        calls.append(prompt)
        return json.dumps({"tasks": [candidate]}, ensure_ascii=False)

    result = generate_project_init_draft([source], [], [], llm_call=llm)

    assert len(calls) == 2
    task = result.tasks[0]
    assert task.goal == "岗位真实运行"
    assert task.acceptance_criteria == "完成验收"
    assert task.process == "试用 → 复盘"
    assert {"建立记录表", "完成复盘"}.issubset({item.title for item in task.subtasks})


def test_strict_models_reject_coerced_ids_and_extra_fields():
    with pytest.raises(ValidationError):
        Evidence(attachment_id="7", file_name="a.txt", location="lines 1", excerpt="x")
    with pytest.raises(ValidationError):
        AgentTask(title="x", subtasks=[], unexpected="do not trust this")


def test_invalid_or_fenced_llm_json_is_rejected_as_business_error():
    def bad_llm(prompt: str, provider: str) -> str:
        return "```json\n{\"tasks\": [{\"title\": 123}]}\n```"

    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft([chunk("x")], [], [], llm_call=bad_llm)


def test_normalizes_redundant_evidence_label_and_null_optional_text():
    payload = raw_task()
    payload["plan_end"] = None
    payload["evidence"][0]["source_label"] = "plan.txt · lines 1-2"
    payload["subtasks"][0]["plan_end"] = None
    payload["subtasks"][0]["evidence"][0]["source_label"] = "plan.txt · lines 1-2"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].plan_end == ""
    assert result.tasks[0].subtasks[0].plan_end == ""
    assert result.tasks[0].evidence[0].source_label == "plan.txt · lines 1-2"


def test_reconciles_end_only_month_as_a_start_only_date():
    payload = raw_task()
    payload["plan_start"] = ""
    payload["plan_end"] = "2026-06"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    assert task.plan_start == "2026-06-01"
    assert task.plan_end == ""


def test_reconciles_end_only_iso_date_as_a_start_only_date():
    payload = raw_task()
    payload["plan_start"] = ""
    payload["plan_end"] = "2026-06-15"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    assert task.plan_start == "2026-06-15"
    assert task.plan_end == ""


def test_normalizes_month_precision_plan_start_to_a_full_date_without_changing_end():
    payload = raw_task()
    payload["plan_start"] = "2026-05"
    payload["plan_end"] = "2026-06"
    payload["subtasks"][0].update({"plan_start": "2026-07", "plan_end": "2026-08"})

    result = generate_project_init_draft(
        [chunk("计划时间 5-6月")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    subtask = task.subtasks[0]
    assert (task.plan_start, task.plan_end) == ("2026-05-01", "2026-06")
    assert (subtask.plan_start, subtask.plan_end) == ("2026-07-01", "2026-08")


def test_preserves_impossible_end_only_iso_date():
    payload = raw_task()
    payload["plan_start"] = ""
    payload["plan_end"] = "2026-02-31"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    assert task.plan_start == ""
    assert task.plan_end == "2026-02-31"


def test_preserves_plan_end_when_plan_start_is_present():
    payload = raw_task()
    payload["plan_start"] = "2026-06-01"
    payload["plan_end"] = "2026-06-15"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    assert task.plan_start == "2026-06-01"
    assert task.plan_end == "2026-06-15"


def test_reconciles_subtask_end_only_month_as_a_start_only_date():
    payload = raw_task()
    payload["subtasks"][0].update({"plan_start": "", "plan_end": "2026-06"})

    result = generate_project_init_draft(
        [chunk("计划时间\n2026-06")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    subtask = result.tasks[0].subtasks[0]
    assert (subtask.plan_start, subtask.plan_end) == ("2026-06-01", "")


def test_reconciles_empty_parent_description_from_first_subtask_completion_standard():
    payload = raw_task()
    payload["description"] = ""
    payload["subtasks"][0]["evaluation_standard"] = "  "
    payload["subtasks"].append({
        "title": "完成验收",
        "evaluation_standard": "交付首版并完成验收",
    })

    result = generate_project_init_draft(
        [chunk("关键成果\n交付首版")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].description == "交付首版并完成验收"


def test_reconciles_parent_description_that_repeats_its_first_subtask_title():
    payload = raw_task()
    payload["description"] = payload["subtasks"][0]["title"]

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].description == ""


def test_reconciles_parent_description_with_normalized_first_subtask_title():
    payload = raw_task()
    payload["description"] = "完成 方案，确认"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].description == ""


def test_parent_description_matching_only_second_subtask_is_preserved():
    payload = raw_task()
    payload["subtasks"].append({"title": "完成上线验收"})
    payload["description"] = "完成上线验收"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].description == "完成上线验收"


@pytest.mark.parametrize("prompt", [_context_prompt([], [], []), _final_merge_prompt([], [])])
def test_prompts_map_project_outline_terms_to_task_fields(prompt: str):
    assert "专项映射到 task title" in prompt
    assert "关键任务映射到 subtasks" in prompt
    assert "关键成果或完成标准映射到父任务 description" in prompt


def test_normalizes_overlong_evidence_excerpt_without_breaking_source_validation():
    excerpt = "实施交付" * 101
    payload = raw_task(evidence=[{
        "attachment_id": 7,
        "file_name": "plan.txt",
        "location": "lines 1-2",
        "excerpt": excerpt,
    }])

    result = generate_project_init_draft(
        [chunk(excerpt)],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].evidence[0].excerpt == excerpt[:300]
    assert len(result.tasks[0].evidence[0].excerpt) == 300


def test_normalizes_string_helper_names_to_a_name_list():
    payload = raw_task()
    payload["subtasks"][0]["helper_names"] = "李四、王五, 赵六\n钱七"

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].subtasks[0].helper_names == ["李四", "王五", "赵六", "钱七"]


def test_normalizes_model_owned_text_fields_to_contract_limits():
    payload = raw_task()
    payload.update(
        {
            "title": "T" * 201,
            "description": "D" * 2001,
            "plan_end": "E" * 51,
        }
    )
    payload["subtasks"][0].update(
        {
            "title": "S" * 201,
            "description": "D" * 2001,
            "plan_end": "E" * 51,
            "evaluation_standard": "V" * 1001,
        }
    )

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    subtask = task.subtasks[0]
    assert len(task.title) == 200
    assert len(task.description) == 2000
    assert len(task.plan_end) == 50
    assert len(subtask.title) == 200
    assert len(subtask.description) == 2000
    assert len(subtask.plan_end) == 50
    assert len(subtask.evaluation_standard) == 1000


def test_ignores_model_supplied_server_owned_fields():
    payload = raw_task()
    payload.update(
        {
            "owner_id": "not-a-person-id",
            "confidence": "certain",
            "merge_status": "definite_duplicate",
            "duplicate_of": "stale-task-id",
            "duplicate_reason": 123,
            "warnings": "not-a-warning-list",
            "source": {"must": "be server generated"},
        }
    )
    payload["subtasks"][0].update(
        {
            "assignee_id": "not-a-person-id",
            "helper_ids": "not-an-id-list",
            "confidence": "certain",
            "merge_status": "definite_duplicate",
            "duplicate_of": "stale-subtask-id",
            "duplicate_reason": 123,
            "warnings": "not-a-warning-list",
            "source": {"must": "be server generated"},
        }
    )

    result = generate_project_init_draft(
        [chunk("实施交付")],
        [{"id": 1, "name": "张三", "is_active": True}],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    task = result.tasks[0]
    subtask = task.subtasks[0]
    assert task.owner_id == 1
    assert subtask.assignee_id == 1
    assert task.merge_status == "new"
    assert subtask.merge_status == "new"
    assert task.duplicate_of is None
    assert subtask.duplicate_of is None
    assert task.warnings == []
    assert subtask.warnings == []
    assert task.source == "plan.txt · lines 1-2"
    assert subtask.source == "plan.txt · lines 1-2"


def test_missing_subtasks_remains_rejected():
    payload = raw_task() | {"subtasks": []}

    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            llm_call=fake_llm({"tasks": [payload]}),
        )


def test_unknown_business_key_remains_rejected():
    payload = raw_task() | {"负责人": "张三"}

    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            llm_call=fake_llm({"tasks": [payload]}),
        )


def test_arbitrary_unknown_business_key_remains_rejected():
    payload = raw_task() | {"not_a_contract_field": "must not be silently ignored"}

    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            llm_call=fake_llm({"tasks": [payload]}),
        )


def test_invalid_draft_exposes_only_safe_validation_field_metadata():
    payload = raw_task()
    payload["evidence"][0]["attachment_id"] = "not-an-integer"

    with pytest.raises(ProjectInitAiInvalidDraft) as error:
        generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            llm_call=fake_llm({"tasks": [payload]}),
        )

    assert error.value.validation_errors == [
        {"path": "tasks[0].evidence[0].attachment_id", "type": "int_type"}
    ]


def test_model_evidence_locator_is_replaced_with_a_canonical_source_excerpt():
    payload = raw_task()
    payload["evidence"][0]["excerpt"] = "模型杜撰的摘录"
    payload["subtasks"][0]["evidence"][0]["excerpt"] = "模型杜撰的摘录"

    result = generate_project_init_draft(
        [chunk("实施交付来源原文")],
        [],
        [],
        llm_call=fake_llm({"tasks": [payload]}),
    )

    assert result.tasks[0].evidence[0].excerpt == "实施交付来源原文"
    assert result.tasks[0].subtasks[0].evidence[0].excerpt == "实施交付来源原文"


def test_invalid_evidence_locator_remains_rejected_with_a_safe_reason_code():
    payload = raw_task()
    payload["evidence"][0]["file_name"] = "not-a-source.txt"

    with pytest.raises(ProjectInitAiInvalidDraft) as error:
        generate_project_init_draft(
            [chunk("实施交付")],
            [],
            [],
            llm_call=fake_llm({"tasks": [payload]}),
        )

    assert error.value.validation_errors == [
        {"path": "tasks[0].evidence[0]", "type": "untraceable_file"}
    ]


def test_empty_llm_result_is_a_safe_business_error():
    with pytest.raises(ProjectInitAiEmptyResult):
        generate_project_init_draft([chunk("没有明确任务")], [], [], llm_call=fake_llm({"tasks": []}))


def test_batches_keep_source_boundaries_and_stay_under_text_limit():
    prompts: list[str] = []

    def llm(prompt: str, provider: str) -> str:
        prompts.append(prompt)
        source_file, source_location = ("a.txt", "lines 1-80") if "a.txt · lines 1-80" in prompt else ("b.txt", "lines 81-160")
        evidence = [{"attachment_id": 7, "file_name": source_file, "location": source_location, "excerpt": "A"}]
        return json.dumps({"tasks": [raw_task(title="批次任务", assignee_name="", evidence=evidence)]}, ensure_ascii=False)

    long_text = "A" * 22_000
    result = generate_project_init_draft(
        [
            chunk(long_text, name="a.txt", location="lines 1-80"),
            chunk(long_text, name="b.txt", location="lines 81-160"),
        ],
        [],
        [],
        llm_call=llm,
    )

    assert len(prompts) >= 2
    assert all("a.txt · lines 1-80" in prompt or "b.txt · lines 81-160" in prompt for prompt in prompts[:2])
    # The source text itself is bounded at 40,000; the prompt envelope adds a
    # handful of literal ASCII `A` characters such as the word `Agent`.
    assert all(prompt.count("A") <= 40_005 for prompt in prompts)
    assert result.tasks


def test_prompt_forbids_side_effects_and_model_invented_ids():
    prompts: list[str] = []

    def llm(prompt: str, provider: str) -> str:
        prompts.append(prompt)
        return json.dumps({"tasks": [raw_task()]}, ensure_ascii=False)

    generate_project_init_draft([chunk("实施交付")], [{"id": 1, "name": "张三", "is_active": True}], [], llm_call=llm)
    assert "不执行数据库、项目或成员修改" in prompts[0]
    assert "不得发明人员、日期或人员 ID" in prompts[0]


def test_person_matching_keeps_punctuation_and_internal_spaces_significant():
    result = generate_project_init_draft(
        [chunk("实施交付 张三")],
        [{"id": 1, "name": "张-三", "is_active": True}],
        [],
        llm_call=fake_llm({"tasks": [raw_task(owner_name="张三", assignee_name="张三")]}),
    )

    task = result.tasks[0]
    assert task.owner_id is None
    assert task.subtasks[0].assignee_id is None
    assert any(warning.code == "person_not_found" for warning in task.warnings)


def test_owner_or_assignee_is_removed_from_helpers_with_conflict_warning():
    result = generate_project_init_draft(
        [chunk("实施交付")],
        [{"id": 1, "name": "张三", "is_active": True}],
        [],
        llm_call=fake_llm({"tasks": [raw_task(helper_names=["张三"])]}),
    )

    subtask = result.tasks[0].subtasks[0]
    assert subtask.assignee_id == 1
    assert subtask.helper_ids == []
    assert any(warning.code == "helper_conflicts_with_assignee" for warning in subtask.warnings)


def test_long_source_evidence_uses_canonical_part_location():
    long_text = "A" * 40_001

    def llm(prompt: str, provider: str) -> str:
        location = "lines 1-2 part 1" if "plan.txt · lines 1-2 part 1" in prompt else "lines 1-2 part 2"
        return json.dumps({"tasks": [raw_task(
            title="长来源任务",
            assignee_name="",
            evidence=[{
                "attachment_id": 7,
                "file_name": "plan.txt",
                "location": location,
                "excerpt": "A",
            }],
        )]}, ensure_ascii=False)

    result = generate_project_init_draft(
        [{"attachment_id": 7, "file_name": "plan.txt", "location": "lines 1-2", "text": long_text}],
        [],
        [],
        llm_call=llm,
    )
    assert result.tasks[0].evidence[0].location == "lines 1-2 part 1"


def test_coarse_worksheet_evidence_is_repaired_to_the_matching_source_row():
    coarse_evidence = [{
        "attachment_id": 7,
        "file_name": "plan.xlsx",
        "location": "'推进表'!A1:J30",
        "excerpt": "",
    }]
    result = generate_project_init_draft(
        [
            {
                "attachment_id": 7,
                "file_name": "plan.xlsx",
                "location": "'推进表'!A2:J2",
                "text": "专项甲 任务甲 交付甲",
            },
            {
                "attachment_id": 7,
                "file_name": "plan.xlsx",
                "location": "'推进表'!A3:J3",
                "text": "专项乙 任务乙 交付乙",
            },
        ],
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(
            title="专项乙",
            assignee_name="",
            evidence=coarse_evidence,
        ) | {"subtasks": [{
            "title": "任务乙",
            "evidence": coarse_evidence,
        }]}]}),
    )

    task = result.tasks[0]
    assert task.evidence[0].location == "'推进表'!A3:J3"
    assert task.subtasks[0].evidence[0].location == "'推进表'!A3:J3"


def test_coarse_worksheet_evidence_without_a_matching_source_title_is_rejected():
    with pytest.raises(ProjectInitAiInvalidDraft):
        generate_project_init_draft(
            [{
                "attachment_id": 7,
                "file_name": "plan.xlsx",
                "location": "'推进表'!A2:J2",
                "text": "专项甲 任务甲 交付甲",
            }],
            [],
            [],
            llm_call=fake_llm({"tasks": [raw_task(
                title="专项乙",
                assignee_name="",
                evidence=[{
                    "attachment_id": 7,
                    "file_name": "plan.xlsx",
                    "location": "'推进表'!A1:J30",
                    "excerpt": "",
                }],
            )]}),
        )


@pytest.mark.parametrize("location", [
    "'推进表'!A0:J30",
    "'推进表'!A1:J1048577",
    "'推进表'!A1:XFE30",
])
def test_out_of_bounds_coarse_worksheet_evidence_is_not_repaired(location):
    with pytest.raises(ProjectInitAiInvalidDraft):
        generate_project_init_draft(
            [{
                "attachment_id": 7,
                "file_name": "plan.xlsx",
                "location": "'推进表'!A2:J2",
                "text": "专项甲 任务甲 交付甲",
            }],
            [],
            [],
            llm_call=fake_llm({"tasks": [raw_task(
                title="专项甲",
                assignee_name="",
                evidence=[{
                    "attachment_id": 7,
                    "file_name": "plan.xlsx",
                    "location": location,
                    "excerpt": "",
                }],
            )]}),
        )


def test_structured_spreadsheet_fallback_uses_traceable_row_data_after_invalid_ai_evidence():
    spreadsheet_row = {
        "attachment_id": 7,
        "file_name": "工作推进表_2026-06-04.xlsx",
        "location": "'工作推进表'!A2:J2",
        "text": (
            "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
            "知识资产AI化\t完成标签体系修订\t知识资产标签框架\t负责人确认可复用\t张三\t李四\t王五、赵六\t2026-06\t进行中\t需协调"
        ),
    }
    result = generate_project_init_draft(
        [spreadsheet_row],
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(
            title="无来源任务",
            assignee_name="",
            evidence=[{
                "attachment_id": 7,
                "file_name": "工作推进表_2026-06-04.xlsx",
                "location": "'工作推进表'!A1:J30",
                "excerpt": "",
            }],
        )]}),
    )

    task = result.tasks[0]
    subtask = task.subtasks[0]
    assert (task.title, task.description, task.owner_name) == ("知识资产AI化", "知识资产标签框架", "张三")
    assert (subtask.title, subtask.assignee_name, subtask.helper_names) == ("完成标签体系修订", "李四", ["王五", "赵六"])
    assert (task.plan_start, task.plan_end) == ("2026-06-01", "")
    assert (subtask.plan_start, subtask.plan_end) == ("2026-06-01", "")
    assert task.evidence[0].location == "'工作推进表'!A2:J2"
    assert subtask.evidence[0].location == "'工作推进表'!A2:J2"


def test_structured_spreadsheet_fallback_accepts_alias_headers_and_merged_workstream_values():
    spreadsheet_rows = [
        {
            "attachment_id": 7,
            "file_name": "非规范推进表.xlsx",
            "location": "'推进表'!A2:J2",
            "text": (
                "重点工作\t任务名称\t目标成果\t验收标准\t统筹负责人\t执行人\t协助人\t开始时间\t状态\t备注\n"
                "知识资产AI化\t制定知识获取计划\t专家清单\t完成访谈计划\t张三\t李四\t王五、赵六\t2026-06-01\t进行中\t先访谈"
            ),
        },
        {
            "attachment_id": 7,
            "file_name": "非规范推进表.xlsx",
            "location": "'推进表'!A3:J3",
            "text": (
                "重点工作\t任务名称\t目标成果\t验收标准\t统筹负责人\t执行人\t协助人\t开始时间\t状态\t备注\n"
                "\t建立知识目录\t目录初稿\t完成目录评审\t\t李四\t吴肖\t2026-06-01\t进行中\t按模板整理"
            ),
        },
    ]
    result = generate_project_init_draft(
        spreadsheet_rows,
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(
            title="无来源任务",
            assignee_name="",
            evidence=[{
                "attachment_id": 7,
                "file_name": "非规范推进表.xlsx",
                "location": "'推进表'!A1:J30",
                "excerpt": "",
            }],
        )]}),
    )

    task = result.tasks[0]
    first, second = task.subtasks
    assert (task.title, task.owner_name, task.plan_start, task.plan_end) == (
        "知识资产AI化", "张三", "2026-06-01", "",
    )
    assert (first.title, first.assignee_name, first.helper_names) == (
        "制定知识获取计划", "李四", ["王五", "赵六"],
    )
    assert (second.title, second.assignee_name, second.helper_names) == (
        "建立知识目录", "李四", ["吴肖"],
    )
    assert second.evaluation_standard == "完成目录评审"
    assert second.description == "按模板整理"


def test_recognized_work_plan_rows_do_not_depend_on_ai_field_interpretation():
    spreadsheet_row = {
        "attachment_id": 7,
        "file_name": "推进表.xlsx",
        "location": "'推进表'!A2:J2",
        "text": (
            "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
            "专项甲\t任务甲\t成果甲\t标准甲\t张三\t李四\t王五\t2026-06-01\t进行中\t备注甲"
        ),
    }
    calls: list[str] = []

    result = generate_project_init_draft(
        [spreadsheet_row],
        [],
        [],
        llm_call=lambda prompt: calls.append(prompt) or '{"tasks":[]}',
    )

    assert calls == []
    assert result.model_name == "structured-spreadsheet"
    assert result.tasks[0].subtasks[0].helper_names == ["王五"]


def test_structured_fallback_does_not_discard_mixed_non_spreadsheet_sources():
    spreadsheet = {
        "attachment_id": 7,
        "file_name": "推进表.xlsx",
        "location": "'推进表'!A1:J2",
        "text": (
            "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
            "表格工作\t表格任务\t表格成果\t表格标准\t张三\t李四\t\t2026-06-01\t进行中\t"
        ),
    }
    text_source = {
        "attachment_id": 8,
        "file_name": "补充说明.txt",
        "location": "第 1 行",
        "text": "补充资料中的工作",
    }
    calls: list[str] = []

    def llm(prompt: str, provider: str) -> str:
        calls.append(prompt)
        return json.dumps({
            "tasks": [raw_task(title="文本工作", evidence=[{
                "attachment_id": 8,
                "file_name": "补充说明.txt",
                "location": "第 1 行",
            }])],
        }, ensure_ascii=False)

    result = generate_project_init_draft([spreadsheet, text_source], [], [], llm_call=llm)

    assert calls
    assert result.tasks[0].title == "文本工作"


def test_structured_rows_merge_workstream_titles_after_normalization():
    rows = [
        {
            "attachment_id": 7,
            "file_name": "推进表.xlsx",
            "location": "'推进表'!A1:J2",
            "text": (
                "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
                "客户成功体系\t任务一\t成果\t标准\t张三\t李四\t\t2026-06-01\t进行中\t"
            ),
        },
        {
            "attachment_id": 7,
            "file_name": "推进表.xlsx",
            "location": "'推进表'!A3:J3",
            "text": (
                "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
                " 客户  成功体系 \t任务二\t\t\t\t王五\t\t\t\t"
            ),
        },
    ]

    result = generate_project_init_draft(rows, [], [], llm_call=lambda _: pytest.fail("AI should not run"))

    assert len(result.tasks) == 1
    assert [item.title for item in result.tasks[0].subtasks] == ["任务一", "任务二"]


def test_merged_alias_header_workbook_is_extracted_end_to_end_without_ai(tmp_path):
    path = tmp_path / "推进表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "推进表"
    sheet.append(["重点工作", "任务名称", "目标成果", "验收标准", "统筹负责人", "执行人", "协助人", "开始时间", "状态", "备注"])
    sheet.append(["知识资产AI化", "制定计划", "专家清单", "完成计划", "张三", "李四", "王五", "2026-06-01", "进行中", "先访谈"])
    sheet.append(["", "建立目录", "目录初稿", "完成评审", "", "李四", "赵六", "2026-06-01", "进行中", "按模板整理"])
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("E2:E3")
    workbook.save(path)

    result = generate_project_init_draft(
        parse_project_init_file(path, path.name),
        [],
        [],
        llm_call=lambda _prompt: (_ for _ in ()).throw(AssertionError("AI should not be called")),
    )

    assert result.model_name == "structured-spreadsheet"
    assert result.tasks[0].title == "知识资产AI化"
    assert [item.title for item in result.tasks[0].subtasks] == ["制定计划", "建立目录"]
    assert result.tasks[0].subtasks[1].helper_names == ["赵六"]


def test_multiline_merged_workplan_uses_structured_import(tmp_path):
    path = tmp_path / "工作推进表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["目标", "重点工作", "评价标准", "序号", "关键任务", "责任人", "计划开始时间", "计划结束时间", "协同人", "完成情况", "备注"])
    sheet.append(["目标一\n目标二", "项目运营系统", "完成标准", 1, "完成系统模块梳理", "吴肖、郭熠彬", "2026-07-03", "2026-07-10", "温会林", "进行中", "保留原文"])
    sheet.append(["", "", "", 2, "导入项目测试", "刘万超", "2026-07-13", "", "", "", ""])
    sheet.merge_cells("B2:B3")
    workbook.save(path)

    people = [
        {"id": 5, "name": "吴肖", "is_active": True, "is_project_member": True},
        {"id": 7, "name": "郭熠彬", "is_active": True, "is_project_member": False},
        {"id": 9, "name": "温会林", "is_active": True, "is_project_member": False},
        {"id": 4, "name": "刘万超", "is_active": True, "is_project_member": False},
    ]
    result = generate_project_init_draft(
        parse_project_init_file(path, path.name),
        people,
        [],
        llm_call=lambda _prompt: (_ for _ in ()).throw(AssertionError("chat analysis must not run")),
    )

    first = result.tasks[0].subtasks[0]
    assert result.model_name == "structured-spreadsheet"
    assert result.tasks[0].description == "目标一 目标二"
    assert (first.assignee_name, first.assignee_id) == ("吴肖", 5)
    assert first.helper_names == ["郭熠彬", "温会林"]
    assert first.helper_ids == [7, 9]
    assert (first.plan_start, first.plan_end) == ("2026-07-03", "2026-07-10")
    assert {warning.code for warning in first.warnings} == {"will_join_project"}


def test_real_workplan_aliases_keep_target_roles_dates_and_evaluation_per_row(tmp_path):
    path = tmp_path / "工作推进表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["AI 项目推进计划"])
    sheet.append([
        "目标",
        "重点工作",
        "评价标准",
        "统筹人",
        "序号",
        "关键任务",
        "责任人",
        "计划开始时间",
        "计划计划结束时间",
        "协同人",
        "完成情况",
        "备注",
    ])
    sheet.append([
        "完成客户成功体系建设",
        "一、建立客户成功体系",
        "服务流程和 SOP 可复用",
        "王经理",
        1,
        "梳理现有客户服务流程",
        "张三、李四",
        "2026-07-01",
        "2026-07-05",
        "王五",
        "进行中",
        "输出流程图",
    ])
    sheet.append([
        "",
        "",
        "评价指标二",
        "",
        2,
        "设计客户成功 SOP",
        "赵六",
        "2026-07-06",
        "2026-07-20",
        "王五",
        "未开始",
        "完成评审",
    ])
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("D3:D4")
    workbook.save(path)

    people = [
        {"id": 1, "name": "张三", "is_active": True, "is_project_member": True},
        {"id": 2, "name": "李四", "is_active": True, "is_project_member": True},
        {"id": 3, "name": "王五", "is_active": True, "is_project_member": True},
        {"id": 4, "name": "赵六", "is_active": True, "is_project_member": True},
        {"id": 5, "name": "王经理", "is_active": True, "is_project_member": True},
    ]
    result = generate_project_init_draft(
        parse_project_init_file(path, path.name),
        people,
        [],
        llm_call=lambda _prompt: (_ for _ in ()).throw(AssertionError("chat analysis must not run")),
    )

    task = result.tasks[0]
    first, second = task.subtasks
    assert result.model_name == "structured-spreadsheet"
    assert (task.title, task.description, task.owner_name) == (
        "一、建立客户成功体系",
        "完成客户成功体系建设",
        "王经理",
    )
    assert (first.title, first.assignee_name, first.assignee_id, first.helper_names, first.helper_ids) == (
        "梳理现有客户服务流程",
        "张三",
        1,
        ["李四", "王五"],
        [2, 3],
    )
    assert (first.plan_start, first.plan_end, first.evaluation_standard, first.status, first.description) == (
        "2026-07-01",
        "2026-07-05",
        "服务流程和 SOP 可复用",
        "进行中",
        "输出流程图",
    )
    assert (second.title, second.assignee_name, second.assignee_id, second.helper_names, second.helper_ids) == (
        "设计客户成功 SOP",
        "赵六",
        4,
        ["王五"],
        [3],
    )
    assert (second.plan_start, second.plan_end, second.evaluation_standard, second.status, second.description) == (
        "2026-07-06",
        "2026-07-20",
        "评价指标二",
        "未开始",
        "完成评审",
    )


def test_structured_spreadsheet_fallback_rejects_non_excel_tabular_text():
    source = {
        "attachment_id": 7,
        "file_name": "meeting-notes.txt",
        "location": "lines 1-2",
        "text": (
            "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
            "知识资产AI化\t完成标签体系修订\t知识资产标签框架\t负责人确认可复用\t张三\t李四\t王五\t2026-06\t进行中\t需协调"
        ),
    }
    with pytest.raises(ProjectInitAiInvalidDraft):
        generate_project_init_draft(
            [source],
            [],
            [],
            llm_call=fake_llm({"tasks": [raw_task(
                title="无来源任务",
                assignee_name="",
                evidence=[{
                    "attachment_id": 7,
                    "file_name": "meeting-notes.txt",
                    "location": "lines 1-99",
                    "excerpt": "",
                }],
            )]}),
        )


def test_structured_spreadsheet_fallback_does_not_infer_year_or_make_backwards_ranges():
    spreadsheet_rows = [
        {
            "attachment_id": 7,
            "file_name": "工作推进表_2026-06-04.xlsx",
            "location": "'工作推进表'!A2:J2",
            "text": (
                "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
                "专项甲\t任务甲\t成果甲\t标准甲\t张三\t李四\t\t4-5月\t进行中\t"
            ),
        },
        {
            "attachment_id": 7,
            "file_name": "工作推进表_2026-06-04.xlsx",
            "location": "'工作推进表'!A3:J3",
            "text": (
                "专项\t关键任务\t关键成果\t完成标准\t统筹人\t负责人\t协同成员\t计划时间\t当前状态\t问题与协调\n"
                "专项乙\t任务乙\t成果乙\t标准乙\t张三\t李四\t\t2026-12-1月\t进行中\t"
            ),
        },
    ]
    result = generate_project_init_draft(
        spreadsheet_rows,
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(
            title="无来源任务",
            assignee_name="",
            evidence=[{
                "attachment_id": 7,
                "file_name": "工作推进表_2026-06-04.xlsx",
                "location": "'工作推进表'!A1:J30",
                "excerpt": "",
            }],
        )]}),
    )

    first, second = result.tasks
    assert (first.plan_start, first.plan_end) == ("", "")
    assert (second.plan_start, second.plan_end) == ("2026-12-01", "")


def test_source_without_attachment_id_accepts_only_none_evidence_id():
    result = generate_project_init_draft(
        [source_chunk("原文片段", name="plan.txt", location="lines 1")],
        [],
        [],
        llm_call=fake_llm({"tasks": [raw_task(evidence=[{
            "attachment_id": None,
            "file_name": "plan.txt",
            "location": "lines 1",
            "excerpt": "原文片段",
        }])]}),
    )
    assert result.tasks[0].evidence[0].attachment_id is None


def test_source_chunk_cannot_be_used_to_forge_an_attachment_id():
    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [source_chunk("原文片段", name="plan.txt", location="lines 1")],
            [],
            [],
            llm_call=fake_llm({"tasks": [raw_task(evidence=[{
                "attachment_id": 7,
                "file_name": "plan.txt",
                "location": "lines 1",
                "excerpt": "原文片段",
            }])]}),
        )


def test_same_title_tasks_across_batches_merge_all_evidence_and_subtasks():
    calls: list[str] = []

    def llm(prompt: str, provider: str) -> str:
        calls.append(prompt)
        if "a.txt · lines 1" in prompt:
            evidence = [{"attachment_id": 7, "file_name": "a.txt", "location": "lines 1", "excerpt": "A"}]
            subtask_title = "方案确认"
        else:
            evidence = [{"attachment_id": 8, "file_name": "b.txt", "location": "lines 2", "excerpt": "B"}]
            subtask_title = "上线确认"
        return json.dumps({"tasks": [raw_task(
            title="同标题任务",
            assignee_name="",
            evidence=evidence,
        ) | {"subtasks": [{
            "title": subtask_title,
            "evidence": evidence,
        }]}]}, ensure_ascii=False)

    result = generate_project_init_draft(
        [
            {"attachment_id": 7, "file_name": "a.txt", "location": "lines 1", "text": "A" * 21_000},
            {"attachment_id": 8, "file_name": "b.txt", "location": "lines 2", "text": "B" * 21_000},
        ],
        [],
        [],
        llm_call=llm,
    )

    assert len(calls) >= 2
    assert "最终合并 Agent" in calls[-1]
    assert len(result.tasks) == 1
    assert {item.file_name for item in result.tasks[0].evidence} == {"a.txt", "b.txt"}
    assert {item.title for item in result.tasks[0].subtasks} == {"方案确认", "上线确认"}


def test_all_task_and_subtask_ids_are_positive_strict_ints_or_none():
    with pytest.raises(ValidationError):
        AgentTask.model_validate({**raw_task(), "owner_id": "1"})
    with pytest.raises(ValidationError):
        AgentTask.model_validate({**raw_task(), "duplicate_of": 0})
    with pytest.raises(ValidationError):
        AgentTask.model_validate({**raw_task(), "subtasks": [{
            "title": "x",
            "assignee_id": -1,
            "helper_ids": ["2"],
        }]})


def test_agent_contract_uses_plan_dates_and_forbids_legacy_deadline():
    task = AgentTask.model_validate(raw_task())
    assert task.plan_start == "2026-09-01"
    assert task.plan_end == "2026-09-30"
    with pytest.raises(ValidationError):
        AgentTask.model_validate({**raw_task(), "deadline": "2026-09-30"})


def test_existing_task_and_every_existing_subtask_id_is_strict_positive_int():
    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [chunk("x")],
            [],
            [{"id": 10, "title": "old", "subtasks": [{"id": "11", "title": "old sub"}]}],
            llm_call=fake_llm({"tasks": [raw_task(owner_name="", assignee_name="")]}),
        )
    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft(
            [chunk("x")],
            [],
            [{"id": 10, "title": "old", "subtasks": [{"id": 0, "title": "old sub"}]}],
            llm_call=fake_llm({"tasks": [raw_task(owner_name="", assignee_name="")]}),
        )


def test_batch_prompt_lists_real_attachment_ids_and_source_labels():
    prompts: list[str] = []

    def llm(prompt: str, provider: str) -> str:
        prompts.append(prompt)
        return json.dumps({"tasks": [raw_task(
            owner_name="",
            assignee_name="",
            evidence=[{"attachment_id": None, "file_name": "plan.txt", "location": "lines 1", "excerpt": "原文"}],
        )]}, ensure_ascii=False)

    generate_project_init_draft(
        [source_chunk("原文", name="plan.txt", location="lines 1")],
        [],
        [],
        llm_call=llm,
    )
    assert '"attachment_id": null' in prompts[0]
    source_label = Evidence(file_name="plan.txt", location="lines 1", excerpt="source").source_label
    assert f'"source_label": {json.dumps(source_label, ensure_ascii=False)}' in prompts[0]
    assert '"location": "lines 1"' in prompts[0]


def test_json_parser_extracts_one_balanced_object_with_nested_strings():
    value = _parse_json_response('说明文字 {"tasks": [], "note": "escaped \\"}\\""} 结束')
    assert value == {"tasks": [], "note": 'escaped "}"'}


def test_json_parser_rejects_ambiguous_multiple_json_values_but_extracts_array():
    with pytest.raises(ProjectInitAiError):
        _parse_json_response('前置 {"tasks": []} 后置 {"tasks": []}')
    assert _parse_json_response('前置 [{"tasks": []}] 后置') == [{"tasks": []}]


def test_array_json_is_rejected_by_the_strict_agent_envelope():
    def array_llm(prompt: str, provider: str) -> str:
        return '前置 [{"tasks": []}] 后置'

    with pytest.raises(ProjectInitAiError):
        generate_project_init_draft([chunk("x")], [], [], llm_call=array_llm)


def test_merge_revalidates_and_caps_deduped_evidence():
    evidence_a = [
        {"attachment_id": 7, "file_name": "a.txt", "location": f"line {i}", "excerpt": "A"}
        for i in range(10)
    ]
    evidence_b = [
        {"attachment_id": 8, "file_name": "b.txt", "location": f"line {i}", "excerpt": "B"}
        for i in range(10)
    ]
    first = AgentTask.model_validate(raw_task(owner_name="", assignee_name="", evidence=evidence_a))
    second = AgentTask.model_validate(raw_task(owner_name="", assignee_name="", evidence=evidence_b))
    merged = _merge_tasks([first, second])
    assert len(merged) == 1
    assert len(merged[0].evidence) <= 10
    assert isinstance(merged[0], AgentTask)
    AgentTask.model_validate(merged[0].model_dump())
