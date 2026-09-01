from __future__ import annotations

import base64
import io
import os
import signal
import subprocess
import threading
import zipfile
import zlib
from dataclasses import FrozenInstanceError
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from app.services import project_init_file_parser as parser
from app.services.project_init_file_parser import (
    ProjectInitFileParseError,
    SourceChunk,
    UnsupportedProjectInitFile,
    parse_project_init_file,
)


REAL_PDF_BASE64 = (
    "JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQg"
    "KG9wZW5zb3VyY2UpCjEgMCBvYmoKPDwKL0YxIDIgMCBSCj4+CmVuZG9iagoyIDAg"
    "b2JqCjw8Ci9CYXNlRm9udCAvSGVsdmV0aWNhIC9FbmNvZGluZyAvV2luQW5zaUVu"
    "Y29kaW5nIC9OYW1lIC9GMSAvU3VidHlwZSAvVHlwZTEgL1R5cGUgL0ZvbnQKPj4K"
    "ZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDcgMCBSIC9NZWRpYUJveCBbIDAg"
    "MCAyMDAgMjAwIF0gL1BhcmVudCA2IDAgUiAvUmVzb3VyY2VzIDw8Ci9Gb250IDEg"
    "MCBSIC9Qcm9jU2V0IFsgL1BERiAvVGV4dCAvSW1hZ2VCIC9JbWFnZUMgL0ltYWdl"
    "SSBdCj4+IC9Sb3RhdGUgMCAvVHJhbnMgPDwKCj4+IAogIC9UeXBlIC9QYWdlCj4+"
    "CmVuZG9iago0IDAgb2JqCjw8Ci9QYWdlTW9kZSAvVXNlTm9uZSAvUGFnZXMgNiAw"
    "IFIgL1R5cGUgL0NhdGFsb2cKPj4KZW5kb2JqCjUgMCBvYmoKPDwKL0F1dGhvciAo"
    "YW5vbnltb3VzKSAvQ3JlYXRpb25EYXRlIChEOjIwMDAwMTAxMDAwMDAwKzAwJzAw"
    "JykgL0NyZWF0b3IgKGFub255bW91cykgL0tleXdvcmRzICgpIC9Nb2REYXRlIChE"
    "OjIwMDAwMTAxMDAwMDAwKzAwJzAwJykgL1Byb2R1Y2VyIChSZXBvcnRMYWIgUERG"
    "IExpYnJhcnkgLSBcKG9wZW5zb3VyY2VcKSkgCiAgL1N1YmplY3QgKHVuc3BlY2lm"
    "aWVkKSAvVGl0bGUgKHVudGl0bGVkKSAvVHJhcHBlZCAvRmFsc2UKPj4KZW5kb2Jq"
    "CjYgMCBvYmoKPDwKL0NvdW50IDEgL0tpZHMgWyAzIDAgUiBdIC9UeXBlIC9QYWdl"
    "cwo+PgplbmRvYmoKNyAwIG9iago8PAovRmlsdGVyIFsgL0FTQ0lJODVEZWNvZGUg"
    "L0ZsYXRlRGVjb2RlIF0gL0xlbmd0aCAxMDYKPj4Kc3RyZWFtCkdhcEEuMGItSCYn"
    "UywvVVZ0YGFqXVt1Y1k5cGM/REc2KUNwKUZfVUMsSzAuSVQtPVI7UjBsP1cpXGh0"
    "dE1NPSNzLEpDWCRJbzxxUCwkOlAtWHM1VWFrMW8rNGBcUSxYKT49K1QmSktUfj5l"
    "bmRzdHJlYW0KZW5kb2JqCnhyZWYKMCA4CjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAw"
    "MDAwMDA2MSAwMDAwMCBuIAowMDAwMDAwMDkyIDAwMDAwIG4gCjAwMDAwMDAxOTkg"
    "MDAwMDAgbiAKMDAwMDAwMDM5MiAwMDAwMCBuIAowMDAwMDAwNDYwIDAwMDAwIG4g"
    "CjAwMDAwMDA3MjEgMDAwMDAgbiAKMDAwMDAwMDc4MCAwMDAwMCBuIAp0cmFpbGVy"
    "Cjw8Ci9JRCAKWzwxYzE3ODE5OGZiZGZhNTFiMjU5OTVkODlkNDEwMjA0Mz48MWMx"
    "NzgxOThmYmRmYTUxYjI1OTk1ZDg5ZDQxMDIwNDM+XQolIFJlcG9ydExhYiBnZW5l"
    "cmF0ZWQgUERGIGRvY3VtZW50IC0tIGRpZ2VzdCAob3BlbnNvdXJjZSkKCi9JbmZv"
    "IDUgMCBSCi9Sb290IDQgMCBSCi9TaXplIDgKPj4Kc3RhcnR4cmVmCjk3NgolJUVP"
    "Rgo="
)

REAL_XLS_ZLIB_BASE64 = (
    "eNrtWD9MU0EY/92jpX9S2tdaTICEvJBYEWEgLiz4kFSZitVFJSZaoIOBFFNJjA6K"
    "YkcTEycNCwkOLigLanTQuDiYYHQwMSFpdXQykYQBen739T0tykATxUDu93K/+953"
    "3919r3ffXe/eLUWLswvNJfyGw6hDWQZQX6UTlALuiwkql1KJbu6nJDV2FAJ+Gsh"
    "6L543vPWpMVTjXYKBx55XxMBnSmdxEYMTuay1jehnHzJC+dBLLDBDmjCa2KsY8wj"
    "zHuZHbPmCuY81t5l7ybYohrBkD3b0OLP4tNHGZWGodp9wnU+s6UYj3qhZfP2OqN"
    "h6cSR/ITO+wwpaPSHMgQZ0IJvL5tV7BA8QBM4QulKprmSyiDgN9hxWpAV8d6P6p"
    "aX126sXIP3qRr2PBmszvX8T+7uGB5iCPM9BUiArWadKQkhmJrOXrISVHs/k1mi9"
    "Fs4S7gWOX6ZpEVTLOYe/uSH8GzgsQsSjiLAc5eAwqf+1h9/ep4bT9jnWTPGSX2"
    "l1n/IDEjdUDaoc5hLjJ3cwH2S+ya22sNzMHCevKE+kGx3h2DTb3OLSBPVziPHB3"
    "l8lt5Nc+HriaWvhi32A5PmB0tX4/Ed7Fm20UY1SffVMo1N0ivv3FJ7Zbi6cRWSZ"
    "uemPBcVvmM63SWf3i+CaEWQxisXXBlsI/rpfFusUZKrPWGUTNWJY5J+gr2obDUJ"
    "DQ0NDQ0NDQ0PjP0E4f8/VmcFTORzwIcDn3OusUyrra5Jdi5OYoGeSDpVHkaM8j"
    "ys1zZ+98Aq3LbHFOu59ocIp6j2PMQyzH2M1z186xInq79lyRfPvhVCt/Zdr8fM"
    "f9/8Ds5TLMw=="
)


def test_source_chunk_exposes_source_label_and_is_frozen():
    chunk = SourceChunk("需求.txt", "第 1-2 行", "第一行\n第二行")

    assert chunk.source_label == "需求.txt · 第 1-2 行"
    with pytest.raises(FrozenInstanceError):
        chunk.text = "changed"


def test_parser_enforces_twenty_five_mib_input_limit(tmp_path, monkeypatch):
    assert parser.MAX_INPUT_BYTES == 25 * 1024 * 1024
    monkeypatch.setattr(parser, "MAX_INPUT_BYTES", 3)
    path = tmp_path / "opaque-storage-key"
    path.write_bytes(b"1234")

    with pytest.raises(ProjectInitFileParseError, match="oversized\\.txt"):
        parse_project_init_file(path, "oversized.txt")


@pytest.mark.parametrize(
    ("limit_name", "limit", "members"),
    [
        ("MAX_ZIP_MEMBERS", 1, [("a.xml", b""), ("b.xml", b"")]),
        ("MAX_ZIP_UNCOMPRESSED_BYTES", 3, [("large.xml", b"1234")]),
        ("MAX_ZIP_COMPRESSION_RATIO", 2, [("dense.xml", b"A" * 1000)]),
    ],
)
def test_ooxml_zip_metadata_limits_are_enforced(
    tmp_path,
    monkeypatch,
    limit_name,
    limit,
    members,
):
    path = tmp_path / "opaque-archive"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for member_name, content in members:
            archive.writestr(member_name, content)
    monkeypatch.setattr(parser, limit_name, limit)

    with pytest.raises(ProjectInitFileParseError, match="unsafe\\.docx"):
        parse_project_init_file(path, "unsafe.docx")


def test_oversized_xls_opens_on_demand_and_releases_resources(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "oversized.xls"
    path.write_bytes(b"xls-placeholder")
    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 1)
    seen_cells = []
    open_calls = []
    released = []

    class FakeSheet:
        name = "Oversized"
        nrows = 100
        ncols = 2

        def cell(self, row, column):
            seen_cells.append((row, column))
            value = "x" if (row, column) in {(0, 0), (0, 1)} else ""
            cell_type = (
                parser.xlrd.XL_CELL_TEXT if value else parser.xlrd.XL_CELL_EMPTY
            )
            return SimpleNamespace(value=value, ctype=cell_type)

    class FakeBook:
        datemode = 0

        def sheets(self):
            return [FakeSheet()]

        def release_resources(self):
            released.append(True)

    def fake_open_workbook(_filename, **kwargs):
        open_calls.append(kwargs)
        return FakeBook()

    monkeypatch.setattr(parser.xlrd, "open_workbook", fake_open_workbook)

    with pytest.raises(ProjectInitFileParseError, match="oversized\\.xls"):
        parse_project_init_file(path, "oversized.xls")

    assert len(seen_cells) < 200
    assert open_calls == [{"on_demand": True}]
    assert released == [True]


def test_pdf_page_limit_is_enforced(tmp_path, monkeypatch):
    path = tmp_path / "opaque-pdf"
    path.write_bytes(b"pdf")
    monkeypatch.setattr(parser, "MAX_PDF_PAGES", 1)
    reader = SimpleNamespace(pages=[SimpleNamespace(), SimpleNamespace()])

    with pytest.raises(ProjectInitFileParseError, match="many-pages\\.pdf"):
        parse_project_init_file(path, "many-pages.pdf", pdf_reader_factory=lambda _path: reader)


def test_pdf_page_text_limit_is_enforced(tmp_path, monkeypatch):
    path = tmp_path / "opaque-pdf"
    path.write_bytes(b"pdf")
    monkeypatch.setattr(parser, "MAX_PDF_PAGE_CHARS", 3)
    page = SimpleNamespace(extract_text=lambda: "1234")

    with pytest.raises(ProjectInitFileParseError, match="verbose\\.pdf"):
        parse_project_init_file(
            path,
            "verbose.pdf",
            pdf_reader_factory=lambda _path: SimpleNamespace(pages=[page]),
        )


def test_total_extracted_character_and_chunk_limits_are_enforced(tmp_path, monkeypatch):
    path = tmp_path / "opaque-pdf"
    path.write_bytes(b"pdf")
    extracted_pages = []

    def page(text):
        return SimpleNamespace(
            extract_text=lambda: extracted_pages.append(text) or text
        )

    pages = [page("abc"), page("def"), page("ghi")]

    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 5)
    with pytest.raises(ProjectInitFileParseError, match="too-much-text\\.pdf"):
        parse_project_init_file(
            path,
            "too-much-text.pdf",
            pdf_reader_factory=lambda _path: SimpleNamespace(pages=pages),
        )
    assert extracted_pages == ["abc", "def"]

    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 100)
    monkeypatch.setattr(parser, "MAX_CHUNKS", 1)
    with pytest.raises(ProjectInitFileParseError, match="too-many-chunks\\.pdf"):
        parse_project_init_file(
            path,
            "too-many-chunks.pdf",
            pdf_reader_factory=lambda _path: SimpleNamespace(pages=pages),
        )


def test_real_pdf_is_parsed_by_pypdf(tmp_path):
    path = tmp_path / "opaque-pdf"
    path.write_bytes(base64.b64decode(REAL_PDF_BASE64))

    chunks = parse_project_init_file(path, "plan.pdf")

    assert chunks == [SourceChunk("plan.pdf", "第 1 页", "Project plan\n")]


def test_txt_utf8_bom_preserves_text_and_eighty_line_ranges(tmp_path):
    lines = [f"第 {index} 行" for index in range(1, 82)]
    path = tmp_path / "upload.txt"
    path.write_bytes("\n".join(lines).encode("utf-8-sig"))

    chunks = parse_project_init_file(path, "原始说明.txt")

    assert chunks == [
        SourceChunk("原始说明.txt", "第 1-80 行", "\n".join(lines[:80])),
        SourceChunk("原始说明.txt", "第 81-81 行", lines[80]),
    ]


def test_txt_falls_back_to_gb18030(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("项目背景\n实施范围".encode("gb18030"))

    assert parse_project_init_file(path, "旧版说明.txt") == [
        SourceChunk("旧版说明.txt", "第 1-2 行", "项目背景\n实施范围")
    ]


def test_dispatches_by_original_name_when_storage_path_has_no_suffix(tmp_path):
    path = tmp_path / "opaque-storage-key"
    path.write_text("项目计划", encoding="utf-8")

    assert parse_project_init_file(path, "plan.txt") == [
        SourceChunk("plan.txt", "第 1-1 行", "项目计划")
    ]


def test_docx_preserves_ordered_paragraph_text_and_twenty_paragraph_ranges(tmp_path):
    path = tmp_path / "plan.docx"
    document = Document()
    paragraphs = [f"段落 {index}" for index in range(1, 22)]
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)

    assert parse_project_init_file(path, "实施方案.docx") == [
        SourceChunk("实施方案.docx", "第 1-20 段", "\n".join(paragraphs[:20])),
        SourceChunk("实施方案.docx", "第 21-21 段", paragraphs[20]),
    ]


def test_docx_extracts_table_only_documents_with_traceable_location(tmp_path):
    path = tmp_path / "table-only.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "部署"
    table.cell(1, 1).text = "张三"
    document.save(path)

    assert parse_project_init_file(path, "表格方案.docx") == [
        SourceChunk(
            "表格方案.docx",
            "第 1 个表格 A1:B2",
            "事项\t负责人\n部署\t张三",
        )
    ]


def test_docx_preserves_mixed_paragraph_and_table_document_order(tmp_path):
    path = tmp_path / "mixed.docx"
    document = Document()
    document.add_paragraph("表格前")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格中"
    document.add_paragraph("表格后")
    document.save(path)

    assert parse_project_init_file(path, "混合方案.docx") == [
        SourceChunk("混合方案.docx", "第 1-1 段", "表格前"),
        SourceChunk("混合方案.docx", "第 1 个表格 A1:A1", "表格中"),
        SourceChunk("混合方案.docx", "第 2-2 段", "表格后"),
    ]


def test_docx_table_cell_limit_is_enforced(tmp_path, monkeypatch):
    path = tmp_path / "wide-table.docx"
    document = Document()
    document.add_table(rows=1, cols=2)
    document.save(path)
    monkeypatch.setattr(parser, "MAX_DOCX_TABLE_CELLS", 1)

    with pytest.raises(ProjectInitFileParseError, match="wide-table\\.docx"):
        parse_project_init_file(path, "wide-table.docx")


def test_oversized_docx_stops_iterating_before_all_blocks_are_materialized(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "oversized.docx"
    document = Document()
    for index in range(100):
        document.add_paragraph(f"paragraph {index}")
    document.save(path)
    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 30)
    seen_blocks = []
    real_document_factory = Document

    class TrackingDocument:
        def __init__(self, source_path):
            self._document = real_document_factory(source_path)

        def iter_inner_content(self):
            for block in self._document.iter_inner_content():
                seen_blocks.append(block)
                yield block

    monkeypatch.setattr(parser, "Document", TrackingDocument)

    with pytest.raises(ProjectInitFileParseError, match="oversized\\.docx"):
        parse_project_init_file(path, "oversized.docx")

    assert len(seen_blocks) < 100


def test_xlsx_uses_actual_non_empty_range_and_skips_empty_worksheets(tmp_path):
    path = tmp_path / "plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "实施计划"
    sheet["B2"] = "负责人"
    sheet["C3"] = "张三"
    workbook.create_sheet("空表")
    workbook.save(path)

    assert parse_project_init_file(path, "计划.xlsx") == [
        SourceChunk("计划.xlsx", "'实施计划'!B2:C3", "负责人\t\n\t张三")
    ]


def test_xlsx_parses_when_storage_path_has_no_extension(tmp_path):
    source_path = tmp_path / "plan.xlsx"
    storage_path = tmp_path / "attachment-storage-key"
    workbook = Workbook()
    workbook.active["A1"] = "项目计划"
    workbook.save(source_path)
    source_path.rename(storage_path)

    assert parse_project_init_file(storage_path, "项目计划.xlsx") == [
        SourceChunk("项目计划.xlsx", "'Sheet'!A1:A1", "项目计划")
    ]


def test_xlsx_falls_back_to_formula_when_cached_value_is_absent(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formula Plan"
    sheet["A1"] = "=1+1"
    workbook.save(path)

    assert parse_project_init_file(path, "公式.xlsx") == [
        SourceChunk("公式.xlsx", "'Formula Plan'!A1:A1", "=1+1")
    ]


def test_xlsx_emits_header_context_for_each_populated_table_row(tmp_path):
    path = tmp_path / "plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["事项", "负责人"])
    sheet.append(["部署", "张三"])
    sheet.append(["验收", "李四"])
    workbook.save(path)

    assert parse_project_init_file(path, "计划.xlsx") == [
        SourceChunk("计划.xlsx", "'Sheet'!A2:B2", "事项\t负责人\n部署\t张三"),
        SourceChunk("计划.xlsx", "'Sheet'!A3:B3", "事项\t负责人\n验收\t李四"),
    ]


def test_xlsx_batches_table_evidence_before_the_chunk_limit(tmp_path, monkeypatch):
    path = tmp_path / "large-plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["事项", "负责人"])
    sheet.append(["部署", "张三"])
    sheet.append(["验收", "李四"])
    workbook.save(path)
    monkeypatch.setattr(parser, "MAX_CHUNKS", 1)

    assert parse_project_init_file(path, "大计划.xlsx") == [
        SourceChunk(
            "大计划.xlsx",
            "'Sheet'!A2:B3",
            "事项\t负责人\n部署\t张三\n验收\t李四",
        )
    ]


def test_xlsx_uses_dense_header_after_a_title_row_for_row_evidence(tmp_path):
    path = tmp_path / "titled-plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "项目计划"
    sheet.append([])
    sheet.append(["事项", "负责人"])
    sheet.append(["部署", "张三"])
    sheet.append(["验收", "李四"])
    workbook.save(path)

    assert parse_project_init_file(path, "标题计划.xlsx") == [
        SourceChunk("标题计划.xlsx", "'Sheet'!A4:B4", "事项\t负责人\n部署\t张三"),
        SourceChunk("标题计划.xlsx", "'Sheet'!A5:B5", "事项\t负责人\n验收\t李四"),
    ]


def test_xlsx_preserves_formula_text_when_cached_result_exists(tmp_path):
    source_path = tmp_path / "formula.xlsx"
    path = tmp_path / "formula-with-cache.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.save(source_path)

    with zipfile.ZipFile(source_path) as source, zipfile.ZipFile(
        path, "w", compression=zipfile.ZIP_DEFLATED
    ) as destination:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(
                    b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>"
                )
            destination.writestr(member, content)

    assert parse_project_init_file(path, "formula.xlsx")[0].text == "=1+1"


def test_xlsx_worksheet_limit_is_enforced(tmp_path, monkeypatch):
    path = tmp_path / "many-sheets.xlsx"
    workbook = Workbook()
    workbook.create_sheet("Second")
    workbook.save(path)
    monkeypatch.setattr(parser, "MAX_WORKSHEETS", 1)

    with pytest.raises(ProjectInitFileParseError, match="many-sheets\\.xlsx"):
        parse_project_init_file(path, "many-sheets.xlsx")


@pytest.mark.parametrize(
    ("limit_name", "populated_cell"),
    [
        ("MAX_WORKSHEET_ROWS", "A2"),
        ("MAX_WORKSHEET_COLUMNS", "B1"),
        ("MAX_WORKSHEET_CELLS", "B1"),
    ],
)
def test_xlsx_dimension_limits_are_enforced(
    tmp_path,
    monkeypatch,
    limit_name,
    populated_cell,
):
    path = tmp_path / f"{limit_name}.xlsx"
    workbook = Workbook()
    workbook.active[populated_cell] = "value"
    workbook.save(path)
    monkeypatch.setattr(parser, limit_name, 1)

    with pytest.raises(ProjectInitFileParseError, match=f"{limit_name}\\.xlsx"):
        parse_project_init_file(path, f"{limit_name}.xlsx")


def test_oversized_xlsx_stops_iterating_before_all_rows_are_materialized(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "oversized.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "x"
    workbook.save(path)
    monkeypatch.setattr(parser, "MAX_EXTRACTED_CHARS", 1)
    seen_rows = []

    class FakeCell:
        def __init__(self, value):
            self.value = value
            self.data_type = "s"

    class FakeSheet:
        title = "Oversized"
        max_row = 100
        max_column = 2

        def iter_rows(self, **_kwargs):
            for row_index in range(self.max_row):
                seen_rows.append(row_index)
                yield (FakeCell("x" if row_index in {0, 1} else None), FakeCell(None))

    class FakeWorkbook:
        worksheets = [FakeSheet()]

        def close(self):
            pass

    monkeypatch.setattr(parser, "load_workbook", lambda *_args, **_kwargs: FakeWorkbook())

    with pytest.raises(ProjectInitFileParseError, match="oversized\\.xlsx"):
        parse_project_init_file(path, "oversized.xlsx")

    assert len(seen_rows) < 200


def test_xlsx_closes_cached_workbook_when_formula_load_fails(tmp_path, monkeypatch):
    path = tmp_path / "broken-second-load.xlsx"
    workbook = Workbook()
    workbook.save(path)
    close_calls = []
    load_calls = 0

    class FakeWorkbook:
        worksheets = []

        def close(self):
            close_calls.append("cached")

    def fake_load_workbook(*_args, **_kwargs):
        nonlocal load_calls
        load_calls += 1
        if load_calls == 1:
            return FakeWorkbook()
        raise RuntimeError("formula workbook load failed")

    monkeypatch.setattr(parser, "load_workbook", fake_load_workbook)

    with pytest.raises(ProjectInitFileParseError, match="broken-second-load\\.xlsx"):
        parse_project_init_file(path, "broken-second-load.xlsx")

    assert close_calls == ["cached"]


@pytest.mark.parametrize("failing_workbook", ["cached", "formula"])
def test_xlsx_closes_both_workbooks_when_one_close_fails(
    tmp_path,
    monkeypatch,
    failing_workbook,
):
    path = tmp_path / "close-failure.xlsx"
    workbook = Workbook()
    workbook.save(path)
    close_calls = []

    class FakeWorkbook:
        worksheets = []

        def __init__(self, name):
            self.name = name

        def close(self):
            close_calls.append(self.name)
            if self.name == failing_workbook:
                raise RuntimeError(f"{self.name} close failed")

    workbooks = iter([FakeWorkbook("cached"), FakeWorkbook("formula")])
    monkeypatch.setattr(
        parser,
        "load_workbook",
        lambda *_args, **_kwargs: next(workbooks),
    )

    with pytest.raises(ProjectInitFileParseError, match="close-failure\\.xlsx"):
        parse_project_init_file(path, "close-failure.xlsx")

    assert set(close_calls) == {"formula", "cached"}


def test_worksheet_chunk_rejects_one_shot_rows_iterables():
    rows = iter(
        [
            [None, "value", "other"],
            [None, None, "last"],
        ]
    )

    with pytest.raises(TypeError, match="row factory"):
        parser._worksheet_chunk("Sheet", rows, "plan.xlsx")


def test_xls_uses_sheet_name_and_actual_non_empty_range(tmp_path, monkeypatch):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"xls-placeholder")

    class FakeSheet:
        name = "旧版计划"
        nrows = 4
        ncols = 4

        def cell_value(self, row, column):
            values = {(1, 1): "事项", (2, 2): 3.0}
            return values.get((row, column), "")

        def cell(self, row, column):
            value = self.cell_value(row, column)
            if isinstance(value, float):
                cell_type = parser.xlrd.XL_CELL_NUMBER
            elif value:
                cell_type = parser.xlrd.XL_CELL_TEXT
            else:
                cell_type = parser.xlrd.XL_CELL_EMPTY
            return SimpleNamespace(value=value, ctype=cell_type)

    class FakeBook:
        datemode = 0

        def sheets(self):
            return [FakeSheet()]

        def release_resources(self):
            pass

    monkeypatch.setattr(
        parser.xlrd,
        "open_workbook",
        lambda filename, **kwargs: FakeBook(),
    )

    assert parse_project_init_file(path, "历史计划.xls") == [
        SourceChunk("历史计划.xls", "'旧版计划'!B2:C3", "事项\t\n\t3")
    ]


def test_real_xls_is_parsed_by_xlrd_and_dates_use_workbook_datemode(tmp_path):
    path = tmp_path / "opaque-xls"
    path.write_bytes(zlib.decompress(base64.b64decode(REAL_XLS_ZLIB_BASE64)))

    assert parse_project_init_file(path, "legacy-plan.xls") == [
        SourceChunk(
            "legacy-plan.xls",
            "'Dates & Plan'!A1:B2",
            "2024-01-02T00:00:00\t\n\tOwner",
        )
    ]


def test_pdf_uses_injected_reader_and_preserves_non_empty_page_numbers(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"not-a-real-pdf")
    received_paths = []

    class FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class FakeReader:
        pages = [FakePage("第一页内容"), FakePage("  \n"), FakePage("第三页内容")]

    def reader_factory(received_path):
        received_paths.append(received_path)
        return FakeReader()

    assert parse_project_init_file(
        path,
        "调研报告.pdf",
        pdf_reader_factory=reader_factory,
    ) == [
        SourceChunk("调研报告.pdf", "第 1 页", "第一页内容"),
        SourceChunk("调研报告.pdf", "第 3 页", "第三页内容"),
    ]
    assert received_paths == [path]


class FakeAntiwordProcess:
    def __init__(self, stdout=b"", stderr=b"", returncode=0, timeout=False):
        self.stdout = io.BytesIO(stdout)
        self.stderr = io.BytesIO(stderr)
        self.returncode = returncode
        self.timeout = timeout
        self.killed = False
        self.wait_calls = []

    def wait(self, timeout=None):
        self.wait_calls.append(timeout)
        if self.timeout and not self.killed:
            raise subprocess.TimeoutExpired(["antiword"], timeout)
        return self.returncode

    def poll(self):
        return None if self.timeout and not self.killed else self.returncode

    def kill(self):
        self.killed = True


def test_doc_invokes_antiword_without_shell_and_with_bounded_timeout(tmp_path, monkeypatch):
    path = tmp_path / "-option-shaped-storage-key"
    path.write_bytes(b"legacy-doc")
    calls = []
    process = FakeAntiwordProcess("第一行\n第二行".encode("utf-8"))

    def fake_popen(args, **kwargs):
        calls.append((args, kwargs))
        return process

    monkeypatch.setattr(parser.subprocess, "Popen", fake_popen)

    assert parse_project_init_file(path, "旧版文档.doc") == [
        SourceChunk("旧版文档.doc", "第 1-2 行", "第一行\n第二行")
    ]
    assert calls[0][0] == ["antiword", "-w", "0", str(path.resolve())]
    assert calls[0][1]["stdout"] is subprocess.PIPE
    assert calls[0][1]["stderr"] is subprocess.PIPE
    assert calls[0][1]["shell"] is False
    if os.name == "nt":
        assert calls[0][1]["creationflags"] & subprocess.CREATE_NEW_PROCESS_GROUP
    else:
        assert calls[0][1]["start_new_session"] is True
    assert process.wait_calls == [30]


def test_doc_uses_saved_pid_as_process_group_id_on_posix(tmp_path, monkeypatch):
    path = tmp_path / "legacy-storage"
    path.write_bytes(b"legacy-doc")
    process = FakeAntiwordProcess(stdout=b"content")
    process.pid = 12345
    captured = []

    expected_identifier = process.pid

    def fail_getpgid(_pid):
        raise AssertionError("antiword process group must not be re-discovered")

    def fake_cleanup(received, readers, *, process_group_id=None):
        captured.append((received, process_group_id))
        for reader in readers:
            reader.join(timeout=parser.ANTIWORD_READER_JOIN_SECONDS)

    monkeypatch.setattr(parser.subprocess, "Popen", lambda *_args, **_kwargs: process)
    monkeypatch.setattr(parser, "_cleanup_antiword_process", fake_cleanup)

    original_os = parser.os
    try:
        parser.os = SimpleNamespace(name="posix", getpgid=fail_getpgid)
        parse_project_init_file(path, "legacy.doc")
    finally:
        parser.os = original_os

    assert captured == [(process, expected_identifier)]


def test_doc_decodes_gb18030_when_strict_utf8_fails(tmp_path, monkeypatch):
    path = tmp_path / "legacy-storage"
    path.write_bytes(b"legacy-doc")
    process = FakeAntiwordProcess("旧版内容".encode("gb18030"))

    def fake_popen(_args, **_kwargs):
        return process

    monkeypatch.setattr(parser.subprocess, "Popen", fake_popen)

    assert parse_project_init_file(path, "legacy.doc") == [
        SourceChunk("legacy.doc", "第 1-1 行", "旧版内容")
    ]


@pytest.mark.parametrize(
    ("limit_name", "stream_name"),
    [
        ("MAX_SUBPROCESS_OUTPUT_BYTES", "stdout"),
        ("MAX_ANTIWORD_STDERR_BYTES", "stderr"),
    ],
)
def test_doc_enforces_antiword_stream_limits(
    tmp_path,
    monkeypatch,
    limit_name,
    stream_name,
):
    path = tmp_path / "legacy-storage"
    path.write_bytes(b"legacy-doc")
    monkeypatch.setattr(parser, limit_name, 3)
    process = FakeAntiwordProcess(
        stdout=b"1234" if stream_name == "stdout" else b"",
        stderr=b"1234" if stream_name == "stderr" else b"",
    )

    def fake_popen(_args, **_kwargs):
        return process

    monkeypatch.setattr(parser.subprocess, "Popen", fake_popen)

    with pytest.raises(ProjectInitFileParseError, match="bounded\\.doc"):
        parse_project_init_file(path, "bounded.doc")
    assert process.killed


def test_doc_rejects_output_invalid_in_utf8_and_gb18030(tmp_path, monkeypatch):
    path = tmp_path / "legacy-storage"
    path.write_bytes(b"legacy-doc")
    process = FakeAntiwordProcess(stdout=b"\xff")

    def fake_popen(_args, **_kwargs):
        return process

    monkeypatch.setattr(parser.subprocess, "Popen", fake_popen)

    with pytest.raises(ProjectInitFileParseError, match="invalid-encoding\\.doc"):
        parse_project_init_file(path, "invalid-encoding.doc")


@pytest.mark.parametrize(
    "failure",
    [
        SimpleNamespace(returncode=2),
        FileNotFoundError("antiword missing"),
        subprocess.TimeoutExpired(["antiword"], 30),
    ],
)
def test_doc_normalizes_antiword_failures(tmp_path, monkeypatch, failure):
    path = tmp_path / "broken.doc"
    path.write_bytes(b"broken")

    def fake_popen(*_args, **_kwargs):
        if isinstance(failure, BaseException):
            raise failure
        return FakeAntiwordProcess(returncode=failure.returncode)

    monkeypatch.setattr(parser.subprocess, "Popen", fake_popen)

    with pytest.raises(ProjectInitFileParseError, match="原始损坏文档\\.doc"):
        parse_project_init_file(path, "原始损坏文档.doc")


def test_doc_timeout_kills_antiword_process(tmp_path, monkeypatch):
    path = tmp_path / "slow.doc"
    path.write_bytes(b"legacy-doc")
    process = FakeAntiwordProcess(timeout=True)
    monkeypatch.setattr(parser.subprocess, "Popen", lambda *_args, **_kwargs: process)

    with pytest.raises(ProjectInitFileParseError, match="slow\\.doc"):
        parse_project_init_file(path, "slow.doc")

    assert process.killed


def test_doc_drains_tail_output_before_closing_antiword_pipes(tmp_path, monkeypatch):
    path = tmp_path / "tail.doc"
    path.write_bytes(b"legacy-doc")

    class TailStream:
        def __init__(self):
            self.read_count = 0
            self.closed = False
            self.tail_ready = threading.Event()

        def read(self, _size):
            if self.closed:
                raise ValueError("stream closed")
            self.read_count += 1
            if self.read_count == 1:
                return b"head\n"
            if self.read_count == 2:
                self.tail_ready.wait(timeout=0.2)
                if self.closed:
                    raise ValueError("stream closed")
                return b"tail\n"
            return b""

        def close(self):
            self.closed = True

    process = FakeAntiwordProcess()
    process.stdout = TailStream()
    threading.Timer(0.01, process.stdout.tail_ready.set).start()
    monkeypatch.setattr(parser.subprocess, "Popen", lambda *_args, **_kwargs: process)

    chunks = parse_project_init_file(path, "tail.doc")

    assert len(chunks) == 1
    assert chunks[0].file_name == "tail.doc"
    assert chunks[0].text == "head\ntail"


def test_antiword_cleanup_terminates_residual_process_tree_when_reader_is_alive(
    monkeypatch,
):
    process = FakeAntiwordProcess()
    process.pid = 12345
    terminated = []
    monkeypatch.setattr(
        parser,
        "_terminate_process_tree",
        lambda received, process_group_id=None: terminated.append(
            (received, process_group_id)
        ),
    )
    joins = []
    reader = SimpleNamespace(
        is_alive=lambda: True,
        join=lambda timeout=None: joins.append(timeout),
    )

    parser._cleanup_antiword_process(process, [reader])

    assert terminated == [(process, None)]
    assert joins == [
        parser.ANTIWORD_READER_JOIN_SECONDS,
        parser.ANTIWORD_READER_JOIN_SECONDS,
    ]


def test_antiword_cleanup_bounds_wait_and_reader_join_when_termination_fails():
    process = FakeAntiwordProcess(timeout=True)
    process.pid = 12345

    def failed_kill():
        raise OSError("termination failed")

    process.kill = failed_kill
    joins = []
    reader = SimpleNamespace(join=lambda timeout=None: joins.append(timeout))

    parser._cleanup_antiword_process(process, [reader])

    assert process.wait_calls == [parser.ANTIWORD_CLEANUP_WAIT_SECONDS]
    assert joins == [parser.ANTIWORD_READER_JOIN_SECONDS]


def test_antiword_cleanup_uses_saved_process_group_after_parent_exit(monkeypatch):
    process = FakeAntiwordProcess()
    process.pid = 12345
    saved_pgid = 23456
    process.poll = lambda: 0
    monkeypatch.setattr(
        parser,
        "_terminate_process_tree",
        lambda received, process_group_id=None: terminated.append(
            (received, process_group_id)
        ),
    )
    terminated = []
    reader = SimpleNamespace(is_alive=lambda: True, join=lambda timeout=None: None)

    parser._cleanup_antiword_process(
        process,
        [reader],
        process_group_id=saved_pgid,
    )

    assert terminated == [(process, saved_pgid)]


@pytest.mark.skipif(os.name == "nt", reason="POSIX process groups are not available")
def test_antiword_termination_targets_the_process_group(monkeypatch):
    calls = []
    process = SimpleNamespace(pid=12345, kill=lambda: calls.append("direct"))
    monkeypatch.setattr(
        parser.os,
        "killpg",
        lambda pgid, sig: calls.append((pgid, sig)),
    )

    parser._terminate_process_tree(process, process_group_id=12346)

    assert calls == [(12346, signal.SIGKILL), "direct"]


def test_txt_chunking_is_lazy_before_global_limits_are_enforced(tmp_path):
    path = tmp_path / "many-lines.txt"
    path.write_text("\n".join(f"line {index}" for index in range(161)), encoding="utf-8")

    chunks = parser._parse_txt(path, "many-lines.txt")

    assert not isinstance(chunks, list)


def test_doc_chunking_is_lazy_before_global_limits_are_enforced(tmp_path, monkeypatch):
    path = tmp_path / "many-lines.doc"
    path.write_bytes(b"legacy-doc")
    process = FakeAntiwordProcess(
        "\n".join(f"line {index}" for index in range(161)).encode("utf-8")
    )
    monkeypatch.setattr(parser.subprocess, "Popen", lambda *_args, **_kwargs: process)

    chunks = parser._parse_doc(path, "many-lines.doc")

    assert not isinstance(chunks, list)


def test_corrupt_third_party_file_error_is_normalized_with_original_filename(tmp_path):
    path = tmp_path / "randomized-storage-name.docx"
    path.write_bytes(b"not-a-zip")

    with pytest.raises(ProjectInitFileParseError, match="客户方案\\.docx"):
        parse_project_init_file(path, "客户方案.docx")


def test_unsupported_extension_raises_domain_error(tmp_path):
    path = tmp_path / "archive.csv"
    path.write_text("a,b", encoding="utf-8")

    with pytest.raises(UnsupportedProjectInitFile, match="archive\\.csv"):
        parse_project_init_file(path, "archive.csv")
