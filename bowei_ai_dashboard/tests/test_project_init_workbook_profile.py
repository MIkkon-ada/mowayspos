from pathlib import Path

from openpyxl import Workbook

from app.services.project_init_workbook_profile import profile_project_init_workbook


def _save_workbook(path: Path, configure) -> Path:
    workbook = Workbook()
    configure(workbook)
    workbook.save(path)
    return path


def test_profile_marks_simple_workbook_as_low_risk(tmp_path):
    path = _save_workbook(
        tmp_path / "simple.xlsx",
        lambda workbook: workbook.active.append(["任务", "负责人"]),
    )

    profile = profile_project_init_workbook(path, "simple.xlsx")

    assert profile["risk_level"] == "low"
    assert profile["signals"] == []
    assert profile["summary"]["visible_sheet_count"] == 1


def test_profile_marks_hidden_merged_formula_workbook_as_high_risk(tmp_path):
    def configure(workbook: Workbook) -> None:
        visible = workbook.active
        visible.append(["阶段", "任务", "负责人"])
        visible.append(["设计", "搭建规则", "张三"])
        visible.merge_cells("A2:A3")
        visible["C3"] = "=COUNTA(B2:B2)"
        hidden = workbook.create_sheet("人员映射")
        hidden.sheet_state = "hidden"
        hidden.append(["姓名", "账号"])
        hidden.append(["张三", "zhangsan"])

    path = _save_workbook(tmp_path / "complex.xlsx", configure)

    profile = profile_project_init_workbook(path, "complex.xlsx")

    assert profile["risk_level"] == "high"
    assert profile["summary"] == {
        "worksheet_count": 2,
        "visible_sheet_count": 1,
        "hidden_sheet_count": 1,
        "merged_range_count": 1,
        "formula_cell_count": 1,
    }
    assert set(profile["signals"]) == {"hidden_sheets", "merged_cells", "formulas"}


def test_profile_does_not_claim_structural_inspection_for_non_xlsx(tmp_path):
    path = tmp_path / "source.txt"
    path.write_text("任务\t负责人", encoding="utf-8")

    profile = profile_project_init_workbook(path, "source.txt")

    assert profile == {
        "risk_level": "low",
        "signals": [],
        "summary": {"inspection": "not_applicable"},
    }
