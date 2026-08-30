from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.project_init_workbook_renderer import (
    WorkbookRenderLimitError,
    render_workbook_images,
)


def _save_workbook(path: Path, configure) -> Path:
    workbook = Workbook()
    configure(workbook)
    workbook.save(path)
    return path


def test_renderer_renders_visible_sheet_and_omits_hidden_sheet(tmp_path):
    def configure(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.title = "推进表"
        sheet.append(["重点工作", "任务"])
        sheet.append(["知识资产AI化", "建立目录"])
        hidden = workbook.create_sheet("内部映射")
        hidden.sheet_state = "hidden"
        hidden.append(["不应渲染"])

    source = _save_workbook(tmp_path / "source.xlsx", configure)
    output = tmp_path / "rendered"

    rendered = render_workbook_images(source, output)

    assert len(rendered) == 1
    assert rendered[0].suffix == ".png"
    assert rendered[0].is_file()


def test_renderer_renders_merged_anchor_value_once(tmp_path):
    def configure(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.append(["重点工作", "任务"])
        sheet.append(["知识资产AI化", "制定计划"])
        sheet.append(["", "建立目录"])
        sheet.merge_cells("A2:A3")

    source = _save_workbook(tmp_path / "merged.xlsx", configure)

    rendered = render_workbook_images(source, tmp_path / "rendered")

    assert rendered[0].stat().st_size > 0


def test_renderer_rejects_workbooks_above_image_limit(tmp_path):
    def configure(workbook: Workbook) -> None:
        workbook.active.append(["任务"])
        workbook.create_sheet("第二页").append(["任务"])

    source = _save_workbook(tmp_path / "many-sheets.xlsx", configure)

    with pytest.raises(WorkbookRenderLimitError, match="image count"):
        render_workbook_images(source, tmp_path / "rendered", max_images=1)
