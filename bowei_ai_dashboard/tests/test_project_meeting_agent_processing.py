from __future__ import annotations

import asyncio
import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import BackgroundTasks, UploadFile
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.database import Base
from app.services.project_meeting_agent import MeetingAgentError, MeetingAgentRunResult
from app.services.project_meeting_agent_contracts import MeetingAgentFinal, MeetingInfo
from app.services.project_meeting_agent_processing import (
    process_project_meeting_agent_run,
    project_meeting_run_status_payload,
)


@pytest.fixture()
def db() -> Session:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _run(db: Session, *, status: str = "queued") -> models.ProjectMeetingRun:
    project = models.Project(id=1, name="Agent project", status="active", is_active=True)
    source = models.MeetingDocumentSource(
        id=2,
        project_id=1,
        original_name="weekly.docx",
        storage_key="1/weekly.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=12,
        content_hash="a" * 64,
    )
    snapshot = {
        "project_id": 1,
        "project": {"id": 1, "name": "Agent project"},
        "members": [],
        "workstreams": [{
            "id": 10,
            "key_task": "Delivery",
            "key_tasks": [{
                "id": 20,
                "title": "Customer list",
                "status": "in_progress",
                "execution_schedules": [{
                    "id": 30,
                    "title": "Customer list first batch",
                    "status": "in_progress",
                    "due_date": "2026-08-20",
                }],
            }],
        }],
        "recent_progress": [],
        "previous_meetings": [],
        "history": {"is_first_meeting": True, "previous_meeting_ids": []},
    }
    run = models.ProjectMeetingRun(
        id=3,
        project_id=1,
        document_source_id=2,
        snapshot_json=json.dumps(snapshot),
        document_text="客户清单已经完成，客户清单第二批需要创建。",
        status=status,
        stage="reading",
    )
    db.add_all([project, source, run])
    db.commit()
    return run


def _final() -> MeetingAgentFinal:
    return MeetingAgentFinal(
        meeting_info=MeetingInfo(
            title="", meeting_date="", meeting_type="", location="", host="",
            participants=[], organizer="", copied_to=[],
        ),
        meeting_info_evidence={}, summary="", summary_evidence=[], agenda_items=[], decisions=[],
        completed_items=[], next_steps=[], risks=[], open_questions=[], task_updates=[],
    )


def _span(document_text: str, quote: str) -> dict[str, int | str]:
    start = document_text.index(quote)
    return {"quote": quote, "char_start": start, "char_end": start + len(quote)}


def _analysis_final(document_text: str, *, action: str = "update_execution_schedule", trusted: bool = True) -> MeetingAgentFinal:
    payload = _final().model_dump(mode="json")
    if action == "update_execution_schedule":
        fact_type, fact_content, field_name, raw_text, value = "completion", "客户清单已经完成", "status", "已经完成", "completed"
        match = {
            "match_id": "M001", "fact_id": "F001", "target_type": "execution_schedule", "target_id": 30,
            "workstream_id": 10, "key_task_id": 20, "confidence": 0.9, "reasons": ["title match"],
            "project_evidence": [{"source_object": "execution_schedule:30", "field": "title", "value": "Customer list first batch"}],
        }
        delta_type = "PROGRESS_UPDATE"
        target = {"project_id": 1, "workstream_id": 10, "key_task_id": 20, "execution_schedule_id": 30}
        before = {"status": "in_progress"}
    else:
        fact_type, fact_content, field_name, raw_text, value = "action_item", "客户清单第二批", "title", "客户清单第二批", "客户清单第二批"
        match = {
            "match_id": "M001", "fact_id": "F001", "target_type": "key_task", "target_id": 20,
            "workstream_id": 10, "key_task_id": 20, "confidence": 0.9, "reasons": ["parent match"],
            "project_evidence": [{"source_object": "key_task:20", "field": "title", "value": "Customer list"}],
        }
        delta_type = "NEW_EXECUTION_SCHEDULE"
        target = {"project_id": 1, "workstream_id": 10, "key_task_id": 20}
        before = {}
    fact = {
        "fact_id": "F001", "fact_type": fact_type, "content": fact_content,
        "fields": {
            field_name: {
                "value": value, "raw_text": raw_text, "evidence": [_span(document_text, raw_text)],
                "provenance": {"source_type": "meeting_fact", "source_fact_id": "F001", "usage": "new"},
            }
        },
        "meeting_evidence": [_span(document_text, fact_content)], "confidence": 0.9, "needs_confirmation": False,
    }
    if not trusted:
        match["target_id"] = 999
    payload.update({
        "meeting_facts": [fact],
        "project_matches": [match],
        "project_deltas": [{"delta_id": "D001", "source_fact_id": "F001", "source_match_id": "M001", "delta_type": delta_type, "reasoning": "baseline comparison"}],
        "proposed_changes": [{
            "change_id": "C001", "source_fact_id": "F001", "source_match_id": "M001", "source_delta_id": "D001",
            "action": action, "target": target, "before": before, "proposed": {field_name: value},
            "field_sources": {field_name: {"source_type": "meeting_fact", "source_fact_id": "F001", "usage": "new"}},
            "requires_confirmation": True,
        }],
    })
    return MeetingAgentFinal.model_validate(payload)


def _agent_result(final: MeetingAgentFinal) -> MeetingAgentRunResult:
    return MeetingAgentRunResult(
        final=final, trace=[{"tool": "search_plan_nodes"}], raw_responses=["{}"],
        invocation_log_ids=[17], events=[{"kind": "final"}], model_code="meeting-model", step_count=2,
    )


def test_meeting_change_proposal_lineage_schema_and_migration_are_declared():
    column = models.MeetingChangeProposal.__table__.c.lineage_json
    assert column.default.arg == "{}"
    assert str(column.server_default.arg) == "{}"
    migration = Path(__file__).parents[1] / "migrations" / "versions" / "d5e6f7a8b9c0_add_meeting_proposal_lineage.py"
    assert migration.exists()
    assert 'down_revision = "c4e5f6a7b8c9"' in migration.read_text(encoding="utf-8")


def test_standard_word_draft_persists_reviewable_work_progress_proposal(db):
    run = _run(db)
    snapshot = json.loads(run.snapshot_json)
    snapshot["standard_minutes"] = {"is_standard_minutes": True}
    run.snapshot_json = json.dumps(snapshot, ensure_ascii=False)
    run.document_text = "客户清单第二批需要创建。"
    db.commit()
    import app.services.project_meeting_agent_processing as processing

    fact_content = "客户清单第二批需要创建"
    normalized = {
        "meeting_draft": {"title": "周会"},
        "next_stage_work": [], "meeting_facts": [{
            "fact_id": "F001", "fact_type": "action_item", "content": fact_content,
            "meeting_evidence": [_span(run.document_text, fact_content)],
            "validation": {"state": "ready", "errors": []},
        }],
        "project_matches": [{
            "match_id": "M001", "fact_id": "F001",
            "target": {"target_type": "key_task", "target_id": 20, "workstream_id": 10, "key_task_id": 20},
            "confidence": 0.91, "reasons": ["客户清单标题一致"],
            "validation": {"state": "ready", "errors": []},
        }],
        "project_deltas": [{
            "delta_id": "D001", "source_fact_id": "F001", "source_match_id": "M001",
            "delta_type": "NEW_EXECUTION_SCHEDULE", "reasoning": "本周新增执行安排",
            "validation": {"state": "ready", "errors": []},
        }],
        "unmatched_items": [], "needs_confirmation": [], "prior_action_tracking": [],
        "proposed_changes": [],
    }
    from app.services.project_meeting_work_progress import build_word_work_progress_preview
    normalized["work_progress_preview"] = build_word_work_progress_preview(normalized, snapshot)
    processing._create_review_draft(db, run, normalized, result_json=json.dumps(normalized, ensure_ascii=False))
    db.commit()

    proposal = db.query(models.MeetingChangeProposal).one()
    assert proposal.action == "confirm_work_progress"
    assert proposal.target_type == "subtask"
    assert proposal.target_id == 20
    assert json.loads(proposal.proposed_json) == {"next_step": fact_content}
    assert json.loads(proposal.lineage_json)["source_fact_ids"] == ["F001"]
    assert db.query(models.KeyTaskExecutionEvent).count() == 0


def test_standard_word_draft_keeps_unlinked_fact_as_explicit_review_item(db):
    run = _run(db)
    snapshot = json.loads(run.snapshot_json)
    snapshot["standard_minutes"] = {"is_standard_minutes": True}
    run.snapshot_json = json.dumps(snapshot, ensure_ascii=False)
    run.document_text = "整理内部知识库"
    db.commit()
    import app.services.project_meeting_agent_processing as processing
    from app.services.project_meeting_work_progress import build_word_work_progress_preview
    normalized = {
        "meeting_draft": {"title": "周会"}, "next_stage_work": [],
        "meeting_facts": [], "project_matches": [], "project_deltas": [],
        "unmatched_items": [], "needs_confirmation": [{
            "fact_id": "F001", "fact_type": "action_item", "content": run.document_text,
            "reason": "有两个可能的关键任务", "meeting_evidence": [_span(run.document_text, run.document_text)],
            "validation": {"state": "ready", "errors": []},
        }],
        "prior_action_tracking": [], "proposed_changes": [],
    }
    normalized["work_progress_preview"] = build_word_work_progress_preview(normalized, snapshot)

    processing._create_review_draft(db, run, normalized, result_json=json.dumps(normalized, ensure_ascii=False))
    db.commit()

    proposal = db.query(models.MeetingChangeProposal).one()
    assert proposal.action == "review_unlinked_fact"
    assert json.loads(proposal.before_json)["fact_id"] == "F001"
    assert json.loads(proposal.proposed_json) == {"next_step": run.document_text}
    assert json.loads(proposal.validation_json)["state"] == "needs_review"


def test_background_success_persists_trusted_update_lineage_and_immutable_result(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing
    monkeypatch.setattr(processing, "run_project_meeting_agent", lambda **_: _agent_result(_analysis_final(run.document_text)))

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    db.expire_all()
    persisted_run = db.get(models.ProjectMeetingRun, run.id)
    change_set = db.query(models.MeetingChangeSet).one()
    proposal = db.query(models.MeetingChangeProposal).one()
    result = json.loads(persisted_run.result_json)
    lineage = json.loads(proposal.lineage_json)
    assert result == json.loads(change_set.result_json)
    assert result["meeting_draft"]["summary"] == ""
    assert result["meeting_facts"][0]["fact_id"] == "F001"
    assert result["proposed_changes"][0]["change_id"] == "C001"
    assert lineage["schema_version"] == 1
    assert lineage["change_id"] == "C001"
    assert lineage["source_fact_id"] == "F001"
    assert lineage["source_match_id"] == "M001"
    assert lineage["source_delta_id"] == "D001"
    assert lineage["field_sources"]["status"]["source_type"] == "meeting_fact"
    assert lineage["meeting_evidence"]["status"][0]["quote"] == "已经完成"
    assert lineage["project_evidence"][0]["source_object"] == "execution_schedule:30"
    assert lineage["delta"]["delta_id"] == "D001"
    assert lineage["baseline_state"] == "existing_target"
    assert lineage["before_baseline"] == {"status": "in_progress"}
    assert lineage["owner_edit_history"] == []
    assert db.query(models.MeetingChangeProposal).count() == 1
    assert db.query(models.ExecutionSchedule).count() == 0


def test_word_meeting_type_is_persisted_without_category_override(db, monkeypatch):
    run = _run(db)
    run.document_text += " Word 原文专题名称"
    db.commit()

    final_payload = _analysis_final(run.document_text).model_dump(mode="json")
    final_payload["meeting_info"]["meeting_type"] = "Word 原文专题名称"
    final_payload["meeting_info_evidence"]["meeting_type"] = [_span(run.document_text, "Word 原文专题名称")]
    final = MeetingAgentFinal.model_validate(final_payload)
    import app.services.project_meeting_agent_processing as processing
    monkeypatch.setattr(processing, "run_project_meeting_agent", lambda **_: _agent_result(final))

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    db.expire_all()
    meeting = db.query(models.Meeting).one()
    persisted_run = db.get(models.ProjectMeetingRun, run.id)
    assert meeting.meeting_type == "Word 原文专题名称"
    assert json.loads(persisted_run.result_json)["meeting_draft"]["meeting_type"] == "Word 原文专题名称"


def test_background_success_persists_create_parent_baseline_without_target_baseline(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing
    monkeypatch.setattr(
        processing,
        "run_project_meeting_agent",
        lambda **_: _agent_result(_analysis_final(run.document_text, action="create_execution_schedule")),
    )

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    proposal = db.query(models.MeetingChangeProposal).one()
    lineage = json.loads(proposal.lineage_json)
    assert lineage["baseline_state"] == "not_applicable_new_object"
    assert lineage["before_baseline"] == {}
    assert lineage["parent_baseline"]["key_task_id"] == 20
    assert lineage["parent_baseline"]["key_task"]["title"] == "Customer list"
    assert lineage["owner_edit_history"] == []


def test_background_does_not_create_proposal_for_untrusted_analysis_change(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing
    monkeypatch.setattr(processing, "run_project_meeting_agent", lambda **_: _agent_result(_analysis_final(run.document_text, trusted=False)))

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    assert db.query(models.MeetingChangeProposal).count() == 0


def test_background_does_not_create_proposal_for_unknown_execution_schedule_field(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing
    payload = _analysis_final(run.document_text).model_dump(mode="json")
    payload["meeting_facts"][0]["fields"] = {
        "lifecycle_status": {
            "value": "已经完成", "raw_text": "已经完成", "evidence": [_span(run.document_text, "已经完成")],
            "provenance": {"source_type": "meeting_fact", "source_fact_id": "F001", "usage": "new"},
        }
    }
    payload["proposed_changes"][0]["proposed"] = {"lifecycle_status": "已经完成"}
    payload["proposed_changes"][0]["field_sources"] = {
        "lifecycle_status": {"source_type": "meeting_fact", "source_fact_id": "F001", "usage": "new"}
    }
    monkeypatch.setattr(
        processing,
        "run_project_meeting_agent",
        lambda **_: _agent_result(MeetingAgentFinal.model_validate(payload)),
    )

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    db.expire_all()
    persisted_run = db.get(models.ProjectMeetingRun, run.id)
    assert json.loads(persisted_run.result_json)["proposed_changes"][0]["validation"]["state"] == "blocked"
    assert db.query(models.MeetingChangeProposal).count() == 0


def test_background_success_persists_agent_audit_then_waits_for_owner_review(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing

    monkeypatch.setattr(
        processing,
        "run_project_meeting_agent",
        lambda **_: MeetingAgentRunResult(
            final=_final(), trace=[{"tool": "get_project_profile"}], raw_responses=["{}"],
            invocation_log_ids=[17], events=[{"kind": "final"}], model_code="meeting-model", step_count=2,
        ),
    )

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    db.expire_all()
    run = db.get(models.ProjectMeetingRun, run.id)
    assert run.status == "pending_review"
    assert run.stage == "pending_review"
    assert run.step_count == 2
    assert run.model_code == "meeting-model"
    assert json.loads(run.invocation_log_ids_json) == [17]
    assert json.loads(run.raw_responses_json) == ["{}"]
    assert json.loads(run.tool_trace_json)["trace"] == [{"tool": "get_project_profile"}]
    assert json.loads(run.result_json)["meeting_draft"]["summary"] == ""
    meeting = db.query(models.Meeting).one()
    assert meeting.source_mode == "ai_analysis"
    assert meeting.review_status == "pending_review"
    assert meeting.publish_status == "draft"
    assert meeting.document_source_id == run.document_source_id
    assert db.query(models.MeetingChangeSet).filter_by(meeting_id=meeting.id).count() == 1


def test_reprocessing_an_unedited_pending_draft_reuses_its_meeting(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing

    monkeypatch.setattr(
        processing,
        "run_project_meeting_agent",
        lambda **_: MeetingAgentRunResult(
            final=_final(), trace=[], raw_responses=["{}"], invocation_log_ids=[28],
            events=[{"kind": "final"}], model_code="meeting-model", step_count=1,
        ),
    )
    factory = sessionmaker(bind=db.get_bind())
    process_project_meeting_agent_run(run.id, session_factory=factory)
    db.expire_all()
    first_meeting = db.query(models.Meeting).one()
    source = db.get(models.MeetingDocumentSource, run.document_source_id)
    retry = models.ProjectMeetingRun(
        id=4,
        project_id=run.project_id,
        document_source_id=run.document_source_id,
        snapshot_json=run.snapshot_json,
        document_text=run.document_text,
        status="queued",
        stage="reading",
    )
    db.add(retry)
    db.commit()

    process_project_meeting_agent_run(retry.id, session_factory=factory)

    db.expire_all()
    assert db.query(models.Meeting).count() == 1
    assert source.meeting_id == first_meeting.id
    assert db.query(models.MeetingChangeSet).filter_by(meeting_id=first_meeting.id).count() == 2


def test_background_agent_error_fails_without_fabricating_result(db, monkeypatch):
    run = _run(db)
    import app.services.project_meeting_agent_processing as processing

    monkeypatch.setattr(
        processing,
        "run_project_meeting_agent",
        lambda **_: (_ for _ in ()).throw(MeetingAgentError(
            "invalid_model_output", "模型未返回有效会议结构", raw_responses=["not-json"],
            trace=[{"tool": "get_project_profile"}], invocation_log_ids=[18], events=[{"kind": "error"}],
        )),
    )

    process_project_meeting_agent_run(run.id, session_factory=sessionmaker(bind=db.get_bind()))

    db.expire_all()
    run = db.get(models.ProjectMeetingRun, run.id)
    assert run.status == "failed"
    assert run.error_code == "invalid_model_output"
    assert "模型未返回有效会议结构" in run.error_message
    assert json.loads(run.result_json) == {}
    assert json.loads(run.raw_responses_json) == ["not-json"]


def test_standard_word_metadata_takes_precedence_over_upload_form_fields():
    from app.services.project_meeting_agent_processing import _apply_meeting_input

    normalized = {"meeting_draft": {
        "title": "Word title", "meeting_date": "2026-08-03", "location": "Word room",
        "host": "Word host", "participants": ["Word participant"],
        "organizer": "Word organizer", "copied_to": ["Word copy"],
    }, "meeting_info_evidence": {"title": {"evidence": []}}}
    original = json.loads(json.dumps(normalized))
    snapshot = {
        "standard_minutes": {"is_standard_minutes": True},
        "meeting_input": {
            "title": "Form title", "meeting_date": "2026-08-04", "location": "Form room",
            "host": "Form host", "participants": "Form participant",
            "organizer": "Form organizer", "copied_to": "Form copy",
        },
    }

    _apply_meeting_input(normalized, snapshot)

    assert normalized == original


def test_status_payload_exposes_stage_steps_error_and_safe_audit_summary(db):
    run = _run(db, status="failed")
    run.stage = "analyzing"
    run.step_count = 3
    run.error_code = "provider_error"
    run.error_message = "模型服务暂不可用"
    run.model_code = "meeting-model"
    run.invocation_log_ids_json = "[8, 9]"
    run.tool_trace_json = json.dumps({"trace": [{"tool": "get_project_profile"}], "events": [{"kind": "model_response"}]})

    payload = project_meeting_run_status_payload(run)

    assert payload["status"] == "failed"
    assert payload["stage"] == "analyzing"
    assert payload["step_count"] == 3
    assert payload["error_code"] == "provider_error"
    assert payload["audit"] == {"model_code": "meeting-model", "invocation_log_ids": [8, 9], "event_count": 1, "tool_call_count": 1}


def test_create_route_queues_background_work_without_calling_model(db, monkeypatch, tmp_path):
    from app.routers import meetings
    import app.services.project_meeting_agent_processing as processing

    db.add_all([
        models.Project(id=1, name="Agent project", status="active", is_active=True),
        models.Person(id=1, name="Owner", is_active=True),
        models.Account(username="owner", password_hash="test", person_id=1),
    ])
    db.commit()
    monkeypatch.setattr(meetings, "require_login", lambda user, _db: user)
    monkeypatch.setattr(meetings, "require_project_access", lambda *_: None)
    monkeypatch.setattr(meetings, "_project_meeting_document_root", lambda: tmp_path)
    monkeypatch.setattr(meetings, "save_meeting_document", lambda *_: {
        "original_name": "weekly.docx", "storage_key": "1/weekly.docx", "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "size_bytes": 10, "content_hash": "b" * 64,
    })
    monkeypatch.setattr(meetings, "extract_meeting_document_text", lambda *_: "会议结论：继续完成验收。")
    monkeypatch.setattr(processing, "process_project_meeting_agent_run", lambda *_: pytest.fail("model worker must not run in create request"))

    upload = UploadFile(filename="weekly.docx", file=BytesIO(b"word"))
    tasks = BackgroundTasks()
    payload = asyncio.run(meetings.create_project_meeting_document_run(
        background_tasks=tasks, project_id=1, file=upload, current_user="owner", db=db,
        title="Weekly meeting", meeting_date="", location="", host="",
        participants="Alice、Bob", organizer="", copied_to="",
    ))

    assert payload["status"] == "queued"
    assert payload["stage"] == "reading"
    assert len(tasks.tasks) == 1
    assert db.query(models.Meeting).count() == 0
    run = db.query(models.ProjectMeetingRun).one()
    assert "requested_meeting_type" not in json.loads(run.snapshot_json)
    assert json.loads(run.snapshot_json)["meeting_input"]["participants"] == "Alice、Bob"


def _standard_minutes_docx_for_review() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Weekly meeting 会议纪要")
    info = document.add_table(rows=0, cols=2)
    for key, value in (
        ("会议时间", "2026-07-27（周一）"),
        ("会议地点", "会议室"),
        ("会议类型", "项目周会"),
        ("会议主持人", "主持人甲"),
        ("与会者", "参会人甲、参会人乙"),
    ):
        cells = info.add_row().cells
        cells[0].text = key
        cells[1].text = value
    document.add_paragraph("一、会议议程")
    document.add_paragraph("1、核对本周计划")
    document.add_paragraph("二、会议小结与决议")
    document.add_paragraph("（一）项目进展")
    document.add_paragraph("· 确认本周计划按节点推进")
    document.add_paragraph("三、待办事项跟踪")
    document.add_paragraph("（一）本周新增待办事项")
    action_table = document.add_table(rows=1, cols=6)
    for cell, value in zip(action_table.rows[0].cells, ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"]):
        cell.text = value
    for cell, value in zip(
        action_table.add_row().cells,
        ["本周-01", "提交计划清单", "负责人甲", "追踪人乙", "2026-08-03 前", "本周新增；按模板留档"],
    ):
        cell.text = value
    document.add_paragraph("整理人 整理人甲")
    document.add_paragraph("抄送 项目组成员甲")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_review_package_projects_stale_standard_minutes_from_word_without_persisting(db, monkeypatch, tmp_path):
    from app.routers import meetings

    run = _run(db, status="pending_review")
    meeting = models.Meeting(
        id=5,
        project_id=1,
        document_source_id=2,
        title="Old AI title",
        meeting_date="2026-07-27",
        host="Old AI host",
        participants="Old AI participant",
        organizer="Old AI organizer",
        copied_to="项目全体成员",
        summary="AI paraphrase that does not exist in Word",
        publish_status="published",
    )
    run.result_json = json.dumps({
        "meeting_draft": {"title": "Old AI title", "summary": "Old AI summary"},
        "agenda_items": [{"content": "Old AI agenda"}],
        "decisions": [{"content": "Old AI decision"}],
        "next_stage_work": [{"content": "Old AI action"}],
    }, ensure_ascii=False)
    db.add_all([meeting, models.MeetingReviewEvent(meeting_id=5, action="published")])
    db.commit()

    source_path = tmp_path / "weekly.docx"
    source_path.write_bytes(_standard_minutes_docx_for_review())
    monkeypatch.setattr(meetings, "_project_meeting_document_root", lambda: tmp_path)
    monkeypatch.setattr(meetings, "download_meeting_document_path", lambda _root, _key: source_path)

    payload = meetings._project_meeting_payload(run, db)

    assert payload["standard_minutes"] is True
    assert payload["meeting"]["copied_to"] == "项目组成员甲"
    assert payload["meeting"]["participants"] == "参会人甲、参会人乙"
    assert payload["meeting"]["summary"] == ""
    assert payload["result"]["agenda_items"][0]["content"] == "核对本周计划"
    assert payload["result"]["decisions"][0]["topic"] == "项目进展"
    assert payload["result"]["decisions"][0]["content"] == "确认本周计划按节点推进"
    assert payload["result"]["next_stage_work"][0]["code"] == "本周-01"
    assert payload["result"]["next_stage_work"][0]["source_note"] == "本周新增；按模板留档"
    assert json.loads(run.result_json)["agenda_items"][0]["content"] == "Old AI agenda"


def test_review_package_keeps_human_edited_standard_meeting_result(db, monkeypatch, tmp_path):
    from app.routers import meetings

    run = _run(db, status="pending_review")
    meeting = models.Meeting(
        id=6,
        project_id=1,
        document_source_id=2,
        title="Human edited title",
        publish_status="draft",
    )
    run.result_json = json.dumps({"meeting_draft": {"title": "Human edited title"}, "agenda_items": [{"content": "Human edited agenda"}]}, ensure_ascii=False)
    db.add_all([meeting, models.MeetingReviewEvent(meeting_id=6, action="edited")])
    db.commit()

    source_path = tmp_path / "weekly.docx"
    source_path.write_bytes(_standard_minutes_docx_for_review())
    monkeypatch.setattr(meetings, "_project_meeting_document_root", lambda: tmp_path)
    monkeypatch.setattr(meetings, "download_meeting_document_path", lambda _root, _key: source_path)

    payload = meetings._project_meeting_payload(run, db)

    assert payload["standard_minutes"] is False
    assert payload["meeting"]["title"] == "Human edited title"
    assert payload["result"]["agenda_items"][0]["content"] == "Human edited agenda"


def test_owner_review_edits_replace_stale_ai_evidence_and_persist_completed_items(db):
    from app import schemas
    from app.routers.meetings import _sync_project_meeting_review_result

    run = _run(db)
    ai_fact = {
        "content": "Original decision", "owner": "", "tracker": "", "due_date": "",
        "evidence": [{"quote": "Original decision", "char_start": 0, "char_end": 17}],
        "validation": {"state": "ready", "errors": []},
    }
    result = {
        "meeting_draft": {"participants": ["Original"], "summary": "Original summary"},
        "meeting_info_evidence": {"participants": {"evidence": ["Original"]}},
        "summary_evidence": {"evidence": ["Original summary"]},
        "agenda_items": [ai_fact], "decisions": [], "completed_items": [],
        "next_stage_work": [], "risks": [],
    }
    run.result_json = json.dumps(result)
    meeting = models.Meeting(
        id=4, project_id=1, document_source_id=2, title="Review", participants="Alice、Bob",
        summary="Corrected summary", review_status="pending_review", publish_status="draft",
    )
    change_set = models.MeetingChangeSet(project_id=1, meeting_id=4, result_json=json.dumps(result))
    db.add_all([meeting, change_set])
    db.commit()

    payload = schemas.MeetingPayload(
        project_id=1,
        agenda_items_json=json.dumps([{**ai_fact, "content": "Corrected decision"}]),
        completed_items_json=json.dumps([{"content": "Confirmed progress"}]),
        decision_items_json="[]", task_list_json="[]", risk_items_json="[]",
    )
    _sync_project_meeting_review_result(db, meeting, payload, {"participants": "Original", "summary": "Original summary"})
    db.commit()
    saved = json.loads(run.result_json)
    assert saved["agenda_items"][0]["content"] == "Corrected decision"
    assert saved["agenda_items"][0]["evidence"] == []
    assert saved["agenda_items"][0]["human_edited"] is True
    assert saved["completed_items"][0]["content"] == "Confirmed progress"
    assert "participants" not in saved["meeting_info_evidence"]
    assert saved["summary_evidence"]["evidence"] == []
    assert json.loads(change_set.result_json) == saved
