from __future__ import annotations

from io import BytesIO

from docx import Document

from app.services.standard_meeting_minutes import parse_standard_meeting_minutes


def _standard_minutes_docx() -> bytes:
    document = Document()
    document.add_paragraph("AI项目落地周会 会议纪要")
    info = document.add_table(rows=0, cols=2)
    for key, value in (
        ("会议时间", "2026-08-10（周一）"),
        ("会议地点", "线下会议室 + 腾讯会议"),
        ("会议类型", "项目例会"),
        ("会议主持人", "温会林"),
        ("与会者", "冯海林、刘万超"),
    ):
        cells = info.add_row().cells
        cells[0].text = key
        cells[1].text = value
    document.add_paragraph("一、会议议程")
    document.add_paragraph("1、市场推广：明确本周渠道落地路径")
    document.add_paragraph("2、项目管理平台：推进计划闭环")
    document.add_paragraph("二、会议小结与决议")
    document.add_paragraph("（一）市场推广：本周启动抖音渠道")
    document.add_paragraph("三、待办事项跟踪")
    document.add_paragraph("（一）本周新增待办事项")
    current = document.add_table(rows=1, cols=6)
    for cell, value in zip(current.rows[0].cells, ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"]):
        cell.text = value
    for values in (
        ["本周-01", "启动抖音渠道", "刘万超", "杨宇帆", "本周内", "本周新增"],
        ["本周-02", "完善会议纪要", "吴肖", "杨宇帆", "本周内", "本周新增"],
    ):
        cells = current.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    document.add_paragraph("（二）上周待办追踪（2026-08-03）")
    prior = document.add_table(rows=1, cols=5)
    for cell, value in zip(prior.rows[0].cells, ["编号", "上周事项", "负责人", "状态", "本周进展/说明"]):
        cell.text = value
    cells = prior.add_row().cells
    for cell, value in zip(cells, ["上周-01", "保留会议录制", "吴肖", "已完成", "本次已录制"]):
        cell.text = value
    document.add_paragraph("整理人 吴肖")
    document.add_paragraph("抄送 AI升级计划项目组全体成员")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_parses_standard_minutes_without_inventing_people_or_progress():
    parsed = parse_standard_meeting_minutes("minutes.docx", _standard_minutes_docx())

    assert parsed["is_standard_minutes"] is True
    assert parsed["title"] == "AI项目落地周会"
    assert parsed["meeting_date"] == "2026-08-10"
    assert parsed["location"] == "线下会议室 + 腾讯会议"
    assert parsed["host"] == "温会林"
    assert parsed["participants"] == "冯海林、刘万超"
    assert parsed["agenda_items"] == ["市场推广：明确本周渠道落地路径", "项目管理平台：推进计划闭环"]
    assert parsed["summary"] == "（一）市场推广：本周启动抖音渠道"
    assert parsed["current_action_items"][0]["负责人"] == "刘万超"
    assert parsed["prior_action_items"][0]["状态"] == "已完成"
    assert parsed["copied_to"] == "AI升级计划项目组全体成员"
    assert "reports" not in parsed


def test_returns_nonstandard_mode_when_required_sections_are_absent():
    document = Document()
    document.add_paragraph("随手记录")
    output = BytesIO()
    document.save(output)

    assert parse_standard_meeting_minutes("notes.docx", output.getvalue())["is_standard_minutes"] is False


def test_accepts_first_meeting_without_prior_action_items():
    document = Document()
    document.add_paragraph("博维管理咨询 AI升级与项目管理周会 会议纪要")
    info = document.add_table(rows=0, cols=2)
    for key, value in (
        ("会议时间", "2026-07-27（周一）"),
        ("会议地点", "线下会议室+腾讯会议（线上）"),
        ("会议类型", "AI升级与项目管理专题（第一次会议）"),
        ("会议主持人", "杨宇帆"),
        ("与会者", "刘万超、邹奇敏、温会林"),
    ):
        cells = info.add_row().cells
        cells[0].text = key
        cells[1].text = value
    document.add_paragraph("一、会议议程")
    document.add_paragraph("1、梳理项目经理工作流程")
    document.add_paragraph("二、会议小结与决议")
    document.add_paragraph("本次会议明确项目推进方式")
    document.add_paragraph("三、待办事项跟踪")
    document.add_paragraph("（一）本周待办事项")
    current = document.add_table(rows=1, cols=6)
    for cell, value in zip(
        current.rows[0].cells,
        ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"],
    ):
        cell.text = value
    for cell, value in zip(
        current.add_row().cells,
        ["本周-01", "完成流程梳理", "温会林", "", "2026-08-03", "本周新增"],
    ):
        cell.text = value
    footer = document.add_table(rows=1, cols=4)
    for cell, value in zip(footer.rows[0].cells, ["整理人", "吴肖", "抄送", "AI升级计划项目组成员"]):
        cell.text = value
    output = BytesIO()
    document.save(output)

    parsed = parse_standard_meeting_minutes("first-meeting.docx", output.getvalue())

    assert parsed["is_standard_minutes"] is True
    assert parsed["meeting_date"] == "2026-07-27"
    assert parsed["location"] == "线下会议室+腾讯会议（线上）"
    assert parsed["host"] == "杨宇帆"
    assert parsed["participants"] == "刘万超、邹奇敏、温会林"
    assert parsed["prior_action_items"] == []


def test_rejects_a_partial_word_file_that_omits_required_standard_sections():
    document = Document()
    document.add_paragraph("AI项目落地周会 会议纪要")
    document.add_paragraph("一、会议议程")
    document.add_paragraph("1、确认本周安排")
    document.add_paragraph("二、会议小结与决议")
    document.add_paragraph("明确推进计划")
    current = document.add_table(rows=1, cols=6)
    for cell, value in zip(current.rows[0].cells, ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "来源/备注"]):
        cell.text = value
    for cell, value in zip(current.add_row().cells, ["本周-01", "确认计划", "吴肖", "杨宇帆", "本周内", "本周新增"]):
        cell.text = value
    output = BytesIO()
    document.save(output)

    assert parse_standard_meeting_minutes("partial.docx", output.getvalue())["is_standard_minutes"] is False
