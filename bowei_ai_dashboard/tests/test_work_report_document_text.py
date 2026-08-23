from io import BytesIO
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pptx import Presentation

from app.services.work_report_document_text import (
    MAX_DOCUMENT_BYTES,
    MAX_EXTRACTED_CHARS,
    WorkReportDocumentTextError,
    extract_work_report_document_text,
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("完成接口联调")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "风险", "等待供应商文档"
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def test_extracts_docx_paragraphs_and_table_cells():
    text = extract_work_report_document_text("weekly.docx", _docx_bytes())

    assert "完成接口联调" in text
    assert "风险\t等待供应商文档" in text


def test_extracts_xlsx_sheet_and_nonempty_cells():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "周报"
    worksheet.append(["本次完成", "联调"])
    out = BytesIO()
    workbook.save(out)

    text = extract_work_report_document_text("weekly.xlsx", out.getvalue())

    assert "[周报]" in text
    assert "本次完成\t联调" in text


def test_extracts_pptx_slide_text_and_table_cells():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.add_textbox(0, 0, 3_000_000, 300_000).text_frame.text = "下周计划"
    table = slide.shapes.add_table(1, 2, 0, 400_000, 3_000_000, 300_000).table
    table.cell(0, 0).text, table.cell(0, 1).text = "任务", "UAT"
    out = BytesIO()
    deck.save(out)

    text = extract_work_report_document_text("weekly.pptx", out.getvalue())

    assert "[第 1 页]" in text
    assert "下周计划" in text
    assert "任务\tUAT" in text


def test_rejects_empty_unsupported_and_oversized_files():
    with pytest.raises(WorkReportDocumentTextError, match="为空"):
        extract_work_report_document_text("weekly.docx", b"")
    with pytest.raises(WorkReportDocumentTextError, match="仅支持"):
        extract_work_report_document_text("weekly.doc", b"legacy")
    with pytest.raises(WorkReportDocumentTextError, match="20 MB"):
        extract_work_report_document_text("weekly.pdf", b"x" * (MAX_DOCUMENT_BYTES + 1))


def test_rejects_pdf_with_too_many_pages_before_extracting_all_pages():
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(width=72, height=72)
    out = BytesIO()
    writer.write(out)

    with pytest.raises(WorkReportDocumentTextError, match="PDF"):
        extract_work_report_document_text("large.pdf", out.getvalue())


def test_stops_pdf_text_extraction_when_one_page_exceeds_the_character_limit():
    class LargeTextPage:
        def extract_text(self, visitor_text=None):
            assert visitor_text is not None
            visitor_text("x" * (MAX_EXTRACTED_CHARS + 1), None, None, None, None)
            return ""

    class LargeTextReader:
        is_encrypted = False
        pages = [LargeTextPage()]

    with patch("app.services.work_report_document_text.PdfReader", return_value=LargeTextReader()):
        with pytest.raises(WorkReportDocumentTextError, match="5000"):
            extract_work_report_document_text("large.pdf", b"%PDF-1.7")


def test_joins_pdf_text_fragments_within_each_page_without_inserting_newlines():
    class FragmentedPage:
        def extract_text(self, visitor_text=None):
            assert visitor_text is not None
            visitor_text("本周完", None, None, None, None)
            visitor_text("成接口联调", None, None, None, None)
            return ""

    class FragmentedReader:
        is_encrypted = False
        pages = [FragmentedPage()]

    with patch("app.services.work_report_document_text.PdfReader", return_value=FragmentedReader()):
        assert extract_work_report_document_text("weekly.pdf", b"%PDF-1.7") == "本周完成接口联调"
